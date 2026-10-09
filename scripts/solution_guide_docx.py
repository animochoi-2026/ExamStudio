"""Ordered solution-only drawings. Main prepares validated common SVG projections."""
from pathlib import Path
from docx.shared import Pt
from word_math import split_math


def compact_step(question,step):
    # Native equation objects cannot reliably wrap inside a narrow table cell.
    return bool(step.get('view')) and len(step['text'])<=180 and not proof_lines(question,step) and all(kind=='text' for kind,_ in split_math(step['text']))


def steps(question):
    g = question.get('solutionGuide')
    if not g or question.get('exportWithoutSolution'):
        return None
    if g.get('version') != 1 or not isinstance(g.get('steps'), list) or not 1 <= len(g['steps']) <= 24:
        raise ValueError('해설 단계 구조를 확인하세요.')
    plain = '\n\n'.join(f"{i+1}. {s['title']}\n{s['text']}" for i, s in enumerate(g['steps']))
    core = question['solution'].split('\n\n[서술형 채점기준')[0].split('\n\n[참고 코멘트]')[0]
    if plain.strip() != core.strip() or not g.get('basis'):
        return None
    paths = question.get('solutionGuideFigurePaths', {})
    for s in g['steps']:
        if s.get('view') and (s['id'] not in paths or not Path(paths[s['id']]).is_file()):
            raise ValueError('해설 보조그림이 준비되지 않았습니다: '+s['id'])
    return g['steps']


def proof_lines(question, step):
    g = question['solutionGuide']; by_id = {r['id']: r for r in g['relations']}
    lines = []
    for rid in (step.get('view') or {}).get('relationIds', []):
        r = by_id[rid]
        if r['kind'] == 'congruence':
            p = r['points']; lines.append('△'+''.join(p[:3])+' ≅ △'+''.join(p[3:])+' ('+r['criterion']+')')
            for i, ref in enumerate(r['dependsOn'], 1):
                premise=by_id[ref];p=premise['points']
                formula = ''.join(p[:2])+' = '+''.join(p[2:]) if premise['kind']=='equalLength' else '∠'+''.join(p[:3])+' = ∠'+''.join(p[3:])
                lines.append(f'{i}) {formula} — '+premise['reason'])
    return lines


def render(document, question, append_rich, set_font, font_pt, font, after):
    ordered = steps(question)
    if ordered is None:
        return None
    paragraphs=[]
    for i, s in enumerate(ordered, 1):
        title=document.add_paragraph(style='Endnote Text');title.paragraph_format.keep_with_next=True
        title.paragraph_format.space_before=Pt(font_pt*.5)
        run=title.add_run(f"{i}. {s['title']}");set_font(run,font_pt,font);run.bold=True
        compact=compact_step(question,s)
        if compact:
            table=document.add_table(rows=1,cols=2);table.autofit=False
            table.columns[0].width=Pt(136);table.columns[1].width=Pt(106)
            left,right=table.rows[0].cells
            left.width=Pt(136);right.width=Pt(106)
            left.paragraphs[0].add_run().add_picture(question['solutionGuideFigurePaths'][s['id']],width=Pt(128))
            empty=right.paragraphs[0];empty._element.getparent().remove(empty._element)
            lines=append_rich(right,s['text'],style='Endnote Text',keep=False,font_pt=font_pt,after=0,font=font,wrap_math=question.get('_guideWrapMath',False))
            paragraphs.extend(lines)
            for p in lines:p.paragraph_format.keep_together=False;p.paragraph_format.widow_control=True
            continue
        if s.get('view'):
            picture=document.add_paragraph(style='Endnote Text');picture.paragraph_format.keep_with_next=True
            # Fits the existing 86 mm answer column; no change to body pagination.
            picture.add_run().add_picture(question['solutionGuideFigurePaths'][s['id']],width=Pt(220))
        lines=append_rich(document,s['text'],style='Endnote Text',keep=False,font_pt=font_pt,after=0,font=font,wrap_math=question.get('_guideWrapMath',False));paragraphs.extend(lines)
        for line in proof_lines(question,s):paragraphs.extend(append_rich(document,line,style='Endnote Text',keep=False,font_pt=font_pt,after=0,font=font))
        for p in paragraphs:p.paragraph_format.keep_together=False;p.paragraph_format.widow_control=True
    core=question['solution'].split('\n\n[서술형 채점기준')[0].split('\n\n[참고 코멘트]')[0]
    appendix=question['solution'][len(core):]
    if appendix:paragraphs.extend(append_rich(document,appendix,style='Endnote Text',keep=False,font_pt=font_pt,after=0,font=font))
    if paragraphs:paragraphs[-1].paragraph_format.space_after=Pt(after)
    return paragraphs
