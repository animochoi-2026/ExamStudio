"""Modern native headings and arbitrary measured rows, without GUI or AI."""
import sys, unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm
import export_docx as d
import paper_form

class ModernTests(unittest.TestCase):
    def test_measured_continuous_rows(self):
        doc=Document();groups=[]
        for i in range(4):
            p=doc.add_paragraph('문항 '+str(i));groups.append([p._p])
        table=paper_form.question_column(doc,groups,600,0,236,[0,40,85,130])
        # Word serializes row heights in twentieths of a point (twips).
        for row,expected in zip(table.rows[:3],[40,45,45]):self.assertAlmostEqual(row.height.mm,expected,delta=.02)
        self.assertTrue(all(row._tr.find(qn('w:trPr')).find(qn('w:cantSplit')) is not None for row in table.rows))

    def test_score_and_no_score_keep_separate_heading_and_body(self):
        positions=[]
        for score in [7,None]:
            doc=Document();d.configure_document(doc,'검증')
            q=dict(id='q',kind='original',sourceId='q',body='원문은 그대로 보존한다.',answer='3',solution='풀이',choices=[],modernMinimal=True,modernSectionStart=True,originalPoints=score,pointsAtBodyEnd=False,workspaceMm=0)
            d.add_question(doc,q,6,{},100,100,0,student=True,show_labels=False)
            texts=[p.text for p in doc.paragraphs]
            section=next(i for i,t in enumerate(texts) if '[서술형]' in t);head=next(i for i,t in enumerate(texts) if t.startswith('06'));body=next(i for i,t in enumerate(texts) if '원문은 그대로' in t)
            self.assertLess(section,head);self.assertLess(head,body);positions.append((section,head,body))
            self.assertEqual('[7점]' in texts[head],score is not None)
            self.assertEqual(doc.paragraphs[head].paragraph_format.line_spacing.pt,20)
            self.assertEqual(doc.paragraphs[head].paragraph_format.space_after.pt,7)
        self.assertEqual(*positions)

if __name__=='__main__':unittest.main()
