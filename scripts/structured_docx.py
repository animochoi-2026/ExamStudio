"""Editable compound question flow. Same ordered nodes as the HTML preview."""
import math
import re
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from word_math import split_math


def split_source_points(value):
    pairs = {'(': ')', '[': ']', '（': '）', '【': '】'}
    math = list(re.finditer(r'\$\$[\s\S]*?\$\$|(?<!\\)\$[^$\n]*?(?<!\\)\$|\\\([\s\S]*?\\\)|\\\[[\s\S]*?\\\]', value))
    count = r'정답\s*\d{1,2}\s*개'
    separator = r'(?:\s*[/／|｜·ㆍ,，;；:：\-–—]\s*|\s+)'
    score = r'(\d{1,2}(?:\.\d{1,2})?)\s*점'
    combined = re.compile(r'([\(\[（【])(\s*'+count+')'+separator+score+r'\s*([\)\]）】])|(^|[ \t])('+count+')'+separator+score+r'(?=\s*$)')

    def preserve_count(match):
        opening, label, _, closing, prefix, bare_label, _ = match.groups()
        if (opening and pairs[opening] != closing) or any(m.start() <= match.start() < m.end() for m in math):
            return match.group()
        return opening+label.rstrip()+closing if opening else prefix+bare_label

    value = combined.sub(preserve_count, value)
    match = re.search(r'\s*[\[(]\s*(\d{1,2}(?:\.\d)?)\s*점\s*[)\]]\s*$', value) or re.search(r'(?<=[?？])\s*[\[(]\s*(\d{1,2}(?:\.\d)?)\s*점(?:\s*[)\]]|(?=\s*(?:$|\$|\\\()))', value)
    return (value[:match.start()]+value[match.end():]).rstrip() if match else value


def for_output(question):
    """Source-only recovery, leaving the saved recognition response unchanged."""
    if question.get('_outputFragment'):
        return question
    leaders = {}
    for match in re.finditer(r'([.…]{2,}[ \t]*)([①-⑳])', question.get('body', '')):
        leaders.setdefault(match[2], []).append(match[1])
    nodes = []
    for node in question['layoutDocument']['nodes']:
        inlines = []
        for i, raw in enumerate(node['inlines']):
            item = dict(raw)
            if raw['origin'] == 'printed':
                if raw['kind'] == 'text' and node['parentId'] is None and node['type'] == 'paragraph':
                    item['text'] = split_source_points(raw['text'])
                elif raw['kind'] == 'reason' and len(leaders.get(raw['text'].strip(), [])) == 1 and not re.search(r'[.…]{2,}\s*$', node['inlines'][i-1]['text'] if i else ''):
                    item['text'] = leaders[raw['text'].strip()][0]+raw['text']
            inlines.append(item)
        nodes.append(dict(node, inlines=inlines))
    return dict(question, layoutDocument=dict(question['layoutDocument'], nodes=nodes))


def inline_math(value):
    """Unwrap one complete legacy math envelope, retaining actual LaTeX bytes."""
    source = value.strip()
    if source.startswith(('$$', '$', '\\(', '\\[')):
        parts = split_math(source)
        if len(parts) != 1 or parts[0][0] not in ('inline', 'display'):
            raise ValueError('구조 수식 토큰에는 한 개의 수식만 기록하세요.')
        return parts[0][1]
    return source


def active(question):
    mode = question.get('layoutMode', 'auto')
    if mode not in ('auto', 'normal', 'structure'):
        raise ValueError('문항 출력 방식 오류')
    layout = question.get('layoutDocument')
    if mode == 'normal':
        return False
    if not layout:
        if mode == 'structure':
            raise ValueError('구조 정보 없음: 수동 구조 교정 또는 AI 재인식 필요')
        return False
    features = layout.get('features', [])
    if not isinstance(features, list) or any(f not in ('multiElementBox','inlineBlanks','multipleFigures','interleavedFlow','proofSteps') for f in features):
        raise ValueError('복합 배치 특징 오류')
    return mode == 'structure' or len(set(features)) >= 2


def tree(layout):
    if layout.get('version') != 1 or not isinstance(layout.get('nodes'), list) or not 0 < len(layout['nodes']) <= 250:
        raise ValueError('복합 배치 데이터 형식 오류')
    roots, seen, ancestors = [], {}, []
    for raw in layout['nodes']:
        n = dict(raw, children=[])
        if n.get('type') not in ('paragraph', 'proofStep', 'box', 'figureGroup', 'figure', 'arrow') or n.get('origin') not in ('printed', 'handwritten', 'uncertain') or n.get('align') not in ('left', 'center', 'right'):
            raise ValueError('복합 배치 요소 형식 오류')
        if not isinstance(n.get('id'), str) or n['id'] in seen or not isinstance(n.get('widthRatio'), (int, float)) or not math.isfinite(n['widthRatio']) or not 0 < n['widthRatio'] <= 1:
            raise ValueError('복합 배치 ID 또는 상대 너비 오류')
        if n.get('parentId') is None:
            ancestors = []
            roots.append(n)
        elif n['parentId'] not in seen or seen[n['parentId']]['type'] not in ('box', 'figureGroup'):
            raise ValueError('복합 배치 포함관계 오류')
        else:
            if n['parentId'] not in ancestors:
                raise ValueError('복합 배치 읽기 순서 오류')
            ancestors = ancestors[:ancestors.index(n['parentId'])+1]
            seen[n['parentId']]['children'].append(n)
        if n['type'] in ('box', 'figureGroup'):
            ancestors.append(n['id'])
        if len(ancestors) > 8:
            raise ValueError('복합 배치 중첩 오류')
        for x in n.get('inlines', []):
            if x.get('kind') not in ('text', 'math', 'blank', 'reason') or x.get('origin') not in ('printed', 'handwritten', 'uncertain') or x.get('align') not in ('left', 'center', 'right') or not isinstance(x.get('text'), str) or not isinstance(x.get('label'), str):
                raise ValueError('복합 배치 인라인 형식 오류')
            if x['kind'] == 'blank' and (x['text'] or x.get('border') not in ('box', 'underline') or not isinstance(x.get('widthEm'), (float, int)) or not math.isfinite(x['widthEm']) or not 1 <= x['widthEm'] <= 30 or any('①' <= c <= '⑳' for c in x['label'])):
                raise ValueError('빈칸 테두리·표지·너비 오류')
        seen[n['id']] = n
    return roots


def figure_path(question, node):
    if node.get('diagram'):
        return question.get('structureFigurePaths', {}).get(node['id'])
    identifier = node.get('figureId')
    if identifier == 'diagram':
        return question.get('diagramPath')
    ids = question.get('materialIds', [])
    return question.get('materialPaths', [])[ids.index(identifier)] if identifier in ids else None


def estimate(question, api):
    question = for_output(question)
    def height(n, ratio=1):
        if n['origin'] != 'printed':
            return 0
        if n['type'] == 'figureGroup':
            total = sum(c['widthRatio'] for c in n['children'] if c['origin'] == 'printed') or 1
            return max((height(c, ratio*c['widthRatio']/total) for c in n['children']), default=0)+8
        if n['type'] == 'box':
            return sum(height(c, ratio*.92) for c in n['children'])+16
        if n['type'] == 'figure':
            size = api['figure_size'](figure_path(question, n))
            if not size:
                raise ValueError('복합 배치 그림 파일 없음: '+n['id'])
            return size[1]*ratio+8
        text = ''.join(('$'+inline_math(x['text'])+'$') if x['kind']=='math' else '\u00a0'*max(1, round(x['widthEm']*2)) if x['kind']=='blank' else x['text'] for x in n['inlines'] if x['origin']=='printed')
        # Only wrapped prose expands in narrower cells; short math rows and
        # arrows keep their line height. Dividing every row by its width made
        # a tiny arrow count as dozens of lines and forced false overflow pages.
        natural = api['estimate_text'](text)
        return max(natural, math.ceil(api['string_width'](text)/(api['column_width']*max(.1, ratio)))*api['line_pt'])+4
    return sum(height(n) for n in tree(question['layoutDocument']))


def render(container, question, api, width, keep=True):
    question = for_output(question)
    missing = next((n for n in question['layoutDocument']['nodes'] if n['type'] == 'arrow' and n['origin'] == 'printed' and not any(x['origin'] == 'printed' and x['kind'] in ('text', 'math') and x['text'].strip() for x in n['inlines'])), None)
    if missing:
        raise ValueError('화살표 원본 확인 필요: '+missing['id']+' — 방향·표지가 없는 구조로 출력할 수 없습니다.')
    w, font_pt = api['w_element'], api['font_pt']
    paragraphs = []
    def table(parent, count, available, border=False):
        t = parent.add_table(rows=1, cols=count)
        t.autofit = False
        props = t._tbl.tblPr
        tw = props.find(qn('w:tblW'))
        tw.set(qn('w:type'), 'pct'); tw.set(qn('w:w'), '5000')
        margins = w('tblCellMar')
        for side in ('top', 'left', 'bottom', 'right'):
            margins.append(w(side, w='70' if border else '0', type='dxa'))
        props.append(margins)
        borders = w('tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            borders.append(w(side, val='single' if border and side not in ('insideH', 'insideV') else 'nil', sz='6', color='444444'))
        props.append(borders)
        if count > 1:
            t.rows[0]._tr.get_or_add_trPr().append(w('cantSplit'))
        for cell in t.rows[0].cells:
            cell.width = Pt(available/count)
            cell._tc.get_or_add_tcPr().append(w('vAlign', val='center'))
            p = cell.paragraphs[0]
            api['readable_paragraph'](p, after=0)
            p.paragraph_format.line_spacing = Pt(1)
            p.paragraph_format.space_after = Pt(0)
        return t
    def paragraph(parent, align):
        p = parent.add_paragraph()
        api['readable_paragraph'](p, after=0 if question.get('_outputFragment') else 3)
        p.alignment = {'left':WD_ALIGN_PARAGRAPH.LEFT, 'center':WD_ALIGN_PARAGRAPH.CENTER, 'right':WD_ALIGN_PARAGRAPH.RIGHT}[align]
        p.paragraph_format.keep_with_next = keep
        p.paragraph_format.keep_together = True
        paragraphs.append(p)
        return p
    def draw(n, parent, available):
        if n['origin'] != 'printed':
            return
        if n['type'] == 'box':
            t = table(parent, 1, available, border=True)
            for c in n['children']:
                draw(c, t.cell(0, 0), available-7)
            return
        if n['type'] == 'figureGroup':
            children = [c for c in n['children'] if c['origin']=='printed']
            if not children:
                return
            t = table(parent, len(children), available)
            total = sum(c['widthRatio'] for c in children)
            for i, c in enumerate(children):
                part = available*c['widthRatio']/total
                t.columns[i].width = Pt(part)
                cell = t.cell(0, i); cell.width = Pt(part)
                cw = cell._tc.get_or_add_tcPr().find(qn('w:tcW'))
                cw.set(qn('w:type'), 'pct'); cw.set(qn('w:w'), str(round(5000*c['widthRatio']/total)))
                draw(c, cell, part)
            return
        if n['type'] == 'figure' and n.get('figureId') in question.get('hiddenFigureIds', []):
            return
        printed = [x for x in n['inlines'] if x['origin'] == 'printed']
        if n['type'] in ('paragraph', 'proofStep') and len(printed) == 1 and printed[0]['kind'] == 'blank' and printed[0]['border'] == 'box':
            # Word trims trailing spaces from a run border. A standalone
            # blank uses a real cell so its label stays centered at every width.
            x = printed[0]
            desired = min(x['widthEm']*font_pt, max(font_pt, available-6))
            t = table(parent, 1, desired, border=True)
            t.alignment = {'left': WD_TABLE_ALIGNMENT.LEFT, 'center': WD_TABLE_ALIGNMENT.CENTER, 'right': WD_TABLE_ALIGNMENT.RIGHT}[n['align']]
            tw = t._tbl.tblPr.find(qn('w:tblW'))
            tw.set(qn('w:w'), str(max(1, round(5000*desired/available))))
            t.columns[0].width = Pt(desired)
            cell = t.cell(0, 0)
            p = cell.paragraphs[0]
            api['readable_paragraph'](p, after=0)
            p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT}[x['align']]
            p.paragraph_format.keep_with_next = keep
            p.paragraph_format.keep_together = True
            api['set_font'](p.add_run(x['label']), font_pt)
            paragraphs.append(p)
            return
        p = paragraph(parent, n['align'])
        if question.get('_outputFragment') and n['type'] in ('paragraph', 'proofStep'):
            # These slices use Chromium's measured 1.5 line height. Native
            # "at least" otherwise adds the font's leading on every text line
            # and pushes already measured fragments onto unintended pages.
            from docx.enum.text import WD_LINE_SPACING
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            line_pt = n.get('lineHeightPt', font_pt*1.5)
            if not isinstance(line_pt, (int, float)) or not math.isfinite(line_pt) or line_pt < font_pt*1.5-0.02:
                raise ValueError('분할 문항의 측정 줄 높이가 잘못되었습니다.')
            p.paragraph_format.line_spacing = Pt(line_pt)
        if n['type'] == 'figure':
            file = figure_path(question, n)
            size = api['figure_size'](file)
            if not size:
                raise ValueError('복합 배치 그림 파일 없음: '+n['id'])
            scale = min(1, max(1, available-3)/size[0])
            p.add_run().add_picture(file, width=Pt(size[0]*scale), height=Pt(size[1]*scale))
            p._p.find('.//'+qn('wp:docPr')).set('descr', '구조 도형 '+n['id'])
            return
        for x in n['inlines']:
            if x['origin'] != 'printed':
                continue
            if x['kind'] == 'math':
                p._p.append(api['formatted_math'](inline_math(x['text']), display=bool(x.get('display')), font_pt=font_pt))
            elif x['kind'] == 'blank':
                # A single bordered run with nonbreaking padding; no overlaid glyphs.
                desired = min(x['widthEm']*font_pt, max(font_pt, available-6))
                label = x['label']; space_width = max(1, api['string_width']('\u00a0'))
                count = max(0, int((desired-api['string_width'](label))/space_width))
                left = 0 if x['align']=='left' else count if x['align']=='right' else count//2
                run = p.add_run('\u00a0'*left+label+'\u00a0'*(count-left))
                api['set_font'](run, font_pt)
                props = run._r.get_or_add_rPr()
                if x['border']=='box':
                    props.append(w('bdr', val='single', sz='6', space='2', color='222222'))
                else:
                    props.append(w('u', val='single'))
            else:
                api['set_font'](p.add_run(x['text']), font_pt)
    roots = tree(question['layoutDocument'])
    if question.get('bodyBorder'):
        container = table(container, 1, width, border=True).cell(0, 0)
        width -= 7
    for n in roots:
        draw(n, container, width)
    if question.get('pointsAtBodyEnd') and question.get('originalPoints') is not None and paragraphs:
        score = str(question['originalPoints']).strip()
        if score:
            api['set_font'](paragraphs[-1].add_run(' ('+(score if '점' in score else score+'점')+')'), font_pt)
    if paragraphs:
        paragraphs[-1].paragraph_format.keep_with_next = False
    return paragraphs
