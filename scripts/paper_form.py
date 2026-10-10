"""Form decorations around the existing editable document body.

The same browser-rendered form layer is embedded in DOCX and HWPX. Only the
form is an image; question paragraphs, tables and native equations are composed
by export_docx/export_hwpx exactly as before.
"""
from copy import deepcopy
from pathlib import Path
import math
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.enum.text import WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT


def pages(snapshot):
    result = snapshot.get('paperFormPages')
    if result is None:
        return None
    if not isinstance(result, list) or not result:
        raise ValueError('시험지 폼의 페이지 정보가 없습니다.')
    if len(result) != len(snapshot.get('settings', {}).get('measuredPages', [])):
        raise ValueError('시험지 폼과 문제 페이지 수가 다릅니다.')
    for index, p in enumerate(result):
        for key in ['topMm', 'leftMm', 'rightMm', 'bottomMm', 'gapMm', 'introMm']:
            v = p.get(key)
            if isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) or not 0 <= v <= 297:
                raise ValueError('시험지 폼 위치가 잘못되었습니다: '+key)
        if p['topMm']+p['bottomMm']+p['introMm'] >= 287 or p['leftMm']+p['rightMm']+p['gapMm'] >= 190:
            raise ValueError('시험지 폼의 문제 영역이 부족합니다.')
        tops=p.get('questionTopsMm')
        if tops is not None:
            columns=snapshot['settings']['measuredPages'][index]
            if not isinstance(tops,list) or len(tops)!=2 or any(not isinstance(col,list) or len(col)!=len(columns[ci]) for ci,col in enumerate(tops)):
                raise ValueError('측정한 문항 위치와 문항 수가 다릅니다.')
            for col in tops:
                if any(isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v) or not 0<=v<297-p['topMm']-p['bottomMm'] for v in col) or any(a>=b for a,b in zip(col,col[1:])):
                    raise ValueError('측정한 문항 위치가 잘못되었습니다.')
        path=Path(p['backgroundPath'])
        if not path.read_bytes().startswith(b'\x89PNG\r\n\x1a\n'):
            raise ValueError('시험지 폼 이미지가 PNG가 아닙니다.')
    return result


def background(paragraph, path, description, logo=None, transparent=False):
    inline=paragraph.add_run().add_picture(path,width=Mm(logo['widthMm'] if logo else 210),height=Mm(logo['heightMm'] if logo else 297))._inline
    anchor=OxmlElement('wp:anchor')
    for k,v in dict(distT='0',distB='0',distL='0',distR='0',simplePos='0',relativeHeight='0',behindDoc='1',locked='1',layoutInCell='0',allowOverlap='1').items():anchor.set(k,v)
    if logo:anchor.set('behindDoc','1' if transparent else '0');anchor.set('relativeHeight','0' if transparent else '251659264');anchor.set('locked','0')
    elif transparent:anchor.set('relativeHeight','1')
    simple=OxmlElement('wp:simplePos');simple.set('x','0');simple.set('y','0');anchor.append(simple)
    for axis in ['H','V']:
        pos=OxmlElement('wp:position'+axis);pos.set('relativeFrom','page');offset=OxmlElement('wp:posOffset');offset.text=str(int(Mm(logo['xMm' if axis=='H' else 'yMm']))) if logo else '0';pos.append(offset);anchor.append(pos)
    anchor.append(deepcopy(inline.find(qn('wp:extent'))));anchor.append(OxmlElement('wp:wrapNone'))
    props=deepcopy(inline.find(qn('wp:docPr')));props.set('name','ExamPaperForm');props.set('descr',description);anchor.append(props)
    for name in ['wp:cNvGraphicFramePr','a:graphic']:
        node=inline.find(qn(name))
        if node is not None:anchor.append(deepcopy(node))
    inline.getparent().replace(inline,anchor)


def configure(section, form=None, title='', title_font_pt=10, title_font='맑은 고딕'):
    section.different_first_page_header_footer=False
    section.header.is_linked_to_previous=False;section.footer.is_linked_to_previous=False
    for story in [section.header,section.footer]:
        for child in list(story._element):story._element.remove(child)
    section.header_distance=Mm(0)
    if form:
        section.top_margin=Mm(form['topMm']);section.bottom_margin=Mm(form['bottomMm'])
        section.left_margin=Mm(form['leftMm']);section.right_margin=Mm(form['rightMm'])
        paragraph=section.header.add_paragraph();paragraph.paragraph_format.space_after=Pt(0)
        # Readers that preserve header drawing order but ignore relativeHeight
        # still need the transparent form's text above its independent logo.
        if form.get('logo') and form.get('backgroundTransparent',False):background(paragraph,form['logo']['path'],'학원 로고',form['logo'],transparent=True)
        background(paragraph,form['backgroundPath'],form.get('description','시험지 폼'),transparent=form.get('backgroundTransparent',False))
        if form.get('logo') and not form.get('backgroundTransparent',False):background(paragraph,form['logo']['path'],'학원 로고',form['logo'],transparent=False)
    else:
        section.top_margin=Mm(20);section.bottom_margin=Mm(18);section.left_margin=Mm(15);section.right_margin=Mm(15)
        section.header_distance=Mm(8)
        p=section.header.add_paragraph();run=p.add_run(title)
        run.font.size=Pt(title_font_pt);run.font.name=title_font
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),title_font)
    columns=section._sectPr.find(qn('w:cols'));columns.set(qn('w:num'),'1' if form else '2');columns.set(qn('w:sep'),'0' if form else '1');columns.set(qn('w:space'),str(round((form['gapMm'] if form else 9)*1440/25.4)))


def spacer(document, mm):
    if mm <= 0:return
    p=document.add_paragraph();f=p.paragraph_format;f.space_before=Pt(0);f.space_after=Pt(0)
    f.line_spacing_rule=WD_LINE_SPACING.EXACTLY;f.line_spacing=Mm(mm);f.keep_with_next=False;f.keep_together=False
    p.add_run(' ')


def layout_table(document, rows, widths):
    table=document.add_table(rows=rows,cols=len(widths))
    table._tbl.set('{urn:examstudio:hwpx-source:1}role','form-layout')
    table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.LEFT
    for column,width in zip(table.columns,widths):column.width=Pt(width)
    for row in table.rows:
        for cell,width in zip(row.cells,widths):
            cell.width=Pt(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','left','bottom','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'0');e.set(qn('w:type'),'dxa');margins.append(e)
    table._tbl.tblPr.append(margins)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'nil');borders.append(e)
    table._tbl.tblPr.append(borders)
    return table


def question_column(document, groups, height_pt, intro_mm, width_pt, tops_mm=None):
    """Native editable cells keep the lower question at the measured browser position.

    Rows have minimum heights, never exact heights: oversized content is not
    clipped. Text, equations, figures and any nested question tables stay native.
    """
    table=layout_table(document,len(groups),[width_pt])
    for i,(row,nodes) in enumerate(zip(table.rows,groups)):
        cell=row.cells[0];cell.width=Pt(width_pt);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
        for child in list(cell._tc):
            if child.tag!=qn('w:tcPr'):cell._tc.remove(child)
        if i==0 and intro_mm:spacer(cell,intro_mm)
        for node in nodes:
            # Match the preview's continuous prose and choice rows; legacy
            # 4/7-point paragraph gaps accumulated outside the measured area.
            for p in ([node] if node.tag==qn('w:p') else node.iter(qn('w:p'))):
                spacing=p.find('w:pPr/w:spacing',namespaces=p.nsmap)
                if spacing is not None and spacing.get(qn('w:after')) in ('80','140'):
                    spacing.set(qn('w:after'),'0')
                keep=p.find('w:pPr/w:keepNext',namespaces=p.nsmap)
                if keep is not None:keep.set(qn('w:val'),'0')
            cell._tc.append(node)
        if not len(cell.paragraphs) or cell._tc[-1].tag!=qn('w:p'):
            tail=cell.add_paragraph();tail.paragraph_format.line_spacing=Pt(1);tail.paragraph_format.space_after=Pt(0)
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i==0 and len(groups)>1:
            row.height=Mm(tops_mm[1]) if tops_mm is not None else Pt((height_pt-intro_mm*72/25.4)/2+intro_mm*72/25.4)
            row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
    return table


def question_page(document, groups, height_pt, form, width_pt):
    table=layout_table(document,1,[width_pt,form['gapMm']*72/25.4,width_pt])
    for ci in range(2):
        cell=table.cell(0,ci*2)
        initial=cell.paragraphs[0]
        if groups[ci]:
            question_column(cell,groups[ci],height_pt,form['introMm'] if ci==0 else 0,width_pt,(form.get('questionTopsMm') or [None,None])[ci])
            cell._tc.remove(initial._p)
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing=Pt(1);p.paragraph_format.space_after=Pt(0)
    return table
