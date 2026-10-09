"""Web measured layout uses the same choice rows in an editable Word table."""
import sys,unittest,tempfile,zipfile
from pathlib import Path
from lxml import etree
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from export_docx import export_document
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}

class WebChoices(unittest.TestCase):
    def render(self,columns):
        q={'id':'q','sourceId':'q','kind':'original','body':'다음 값을 구하시오.','choices':['$x^2$','2','3','4','5'],'answer':'1','solution':'식을 계산한다.','workspaceMm':15}
        if columns is not None:q['webChoiceColumns']=columns
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/'paper.docx'
            export_document({'title':'확인','settings':{'bodyFontSize':10,'workspaceLines':0,'answerMode':'quick','measuredPages':[[['q'],[]]],'quadrantLayout':False},'questions':[q]},p)
            return etree.fromstring(zipfile.ZipFile(p).read('word/document.xml'))
    def test_two_columns_preserve_math_order_unsplit_rows_and_no_border(self):
        r=self.render(2);tables=r.xpath('/w:document/w:body/w:tbl',namespaces=NS)
        self.assertEqual(len(tables),1)
        self.assertEqual(len(tables[0].xpath('w:tr',namespaces=NS)),3)
        self.assertEqual(len(tables[0].xpath('.//m:oMath',namespaces=NS)),1)
        cells=tables[0].xpath('w:tr/w:tc',namespaces=NS)
        self.assertEqual([''.join(c.itertext()).strip()[:1] for c in cells],['①','②','③','④','⑤',''])
        self.assertEqual(len(tables[0].xpath('w:tr/w:trPr/w:cantSplit',namespaces=NS)),3)
        self.assertEqual(tables[0].xpath('w:tblPr/w:tblBorders/w:top/@w:val',namespaces=NS),['nil'])
        section_paragraph=r.xpath('/w:document/w:body/w:p[w:pPr/w:sectPr]',namespaces=NS)[0]
        self.assertEqual(section_paragraph.xpath('w:pPr/w:spacing/@w:line',namespaces=NS),['20'])
    def test_ordinary_and_vertical_paths_still_use_paragraph_choices(self):
        for mode in (None,1):
            r=self.render(mode)
            self.assertEqual(len(r.xpath('/w:document/w:body/w:tbl',namespaces=NS)),0)
            self.assertEqual(len(r.xpath('//m:oMath',namespaces=NS)),1)

    def test_measured_column_break_has_own_minimal_paragraph(self):
        from docx import Document
        from export_docx import add_question,configure_document
        d=Document();configure_document(d,'Column test')
        q={'id':'q','sourceId':'q','kind':'original','body':'Question','choices':[]}
        add_question(d,q,1,{},50,50,0,new_column=True,separate_column_break=True)
        first=d.paragraphs[0]
        self.assertEqual(first._p.xpath('w:pPr/w:spacing/@w:line'),['20'])
        self.assertEqual(first._p.xpath('w:r/w:br/@w:type'),['column'])
        self.assertEqual(first.text,'')

    def test_terminal_workspace_uses_spacing_without_spill_paragraph(self):
        from docx import Document
        from docx.shared import Pt
        from export_docx import add_question,configure_document
        for size in (7,10,15):
            documents=[]
            for terminal in (False,True):
                d=Document();configure_document(d,'Workspace test')
                d.styles['Normal'].font.size=Pt(size)
                q={'id':'q','sourceId':'q','kind':'original','body':'Question','choices':[],'workspaceMm':15}
                last=add_question(d,q,1,{},50,50,0,terminal_workspace=terminal)
                documents.append(d)
                if terminal:
                    self.assertTrue(last.text)
                    self.assertGreaterEqual(last.paragraph_format.space_after.pt,15*72/25.4-.05)
                    self.assertFalse(last.paragraph_format.keep_with_next)
                else:
                    self.assertEqual(last.text,'')
            self.assertEqual(len(documents[0].paragraphs),len(documents[1].paragraphs)+1)

if __name__=='__main__':unittest.main()
