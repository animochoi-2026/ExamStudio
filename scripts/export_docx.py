"""Export the app's confirmed questions to a two-column A4 DOCX.

CLI: export_docx.py --input snapshot.json --output exam.docx
Stdout is one JSON result. Errors are on stderr and return a nonzero exit code.
Word/Hancom are the final pagination engines: the included manifest explicitly
reports *estimates*, not a claim that DOCX has already been visually rendered.
"""
from __future__ import annotations

import argparse
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
from copy import deepcopy
from functools import lru_cache
from io import BytesIO

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree
from PIL import Image, ImageFont

from word_math import MathSyntaxError, dimensions, latex_to_omml, parse_latex, split_math, text_width_em

BODY_FONT = "맑은 고딕"
FONT_PT = 12
NOTE_FONT = "맑은 고딕"
NOTE_FONT_PT = 9
LINE_PT = 18
NOTE_LINE_PT = 13.5
COL_WIDTH_PT = 85.5 * 72 / 25.4
COL_HEIGHT_PT = 259 * 72 / 25.4
HALF_HEIGHT_PT = (COL_HEIGHT_PT - 18) / 2
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
PPR_ORDER = "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange".split()
RPR_ORDER = "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath rPrChange".split()
SECT_ORDER = "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings sectPrChange".split()


def canonicalize_properties(root):
    """Order CT_PPr/CT_RPr/CT_SectPr children exactly as OOXML requires.

    Manual properties appended out of order can also confuse python-docx's
    later successor lookup. Normalize every generated part before serialization.
    """
    for tag, order in (("pPr", PPR_ORDER), ("rPr", RPR_ORDER), ("sectPr", SECT_ORDER)):
        ranks = {qn("w:"+name): index for index, name in enumerate(order)}
        for parent in root.iter(qn("w:"+tag)):
            children = list(parent)
            for child in sorted(children, key=lambda child: ranks.get(child.tag, len(ranks))):
                parent.append(child)


def set_font(style, size=None, font=None):
    font = BODY_FONT if font is None else font
    size = FONT_PT if size is None else size
    style.font.name = font
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        if key.lower().endswith("theme"):
            del fonts.attrib[key]
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), font)
    properties = style.element.get_or_add_rPr()
    complex_size = properties.find(qn("w:szCs"))
    if complex_size is None:
        complex_size = w_element("szCs")
        properties.append(complex_size)
    complex_size.set(qn("w:val"), str(round(size*2)))


def w_element(tag, **attrs):
    elem = OxmlElement("w:" + tag)
    for key, value in attrs.items():
        elem.set(qn("w:" + key), str(value))
    return elem


def font_properties(font_pt, font=None):
    font = BODY_FONT if font is None else font
    properties = w_element("rPr")
    properties.extend([w_element("rFonts", ascii=font, hAnsi=font, eastAsia=font, cs=font),
                       w_element("sz", val=round(font_pt*2)), w_element("szCs", val=round(font_pt*2))])
    return properties


def formatted_math(source, display=False, font_pt=None):
    font_pt = FONT_PT if font_pt is None else font_pt
    """Native OMML with explicit body/note size, including fraction controls."""
    equation = latex_to_omml(source, display=display)
    for run in equation.iter(qn("m:r")):
        properties = font_properties(font_pt, "Cambria Math")
        text = run.find(qn("m:t"))
        run.insert(list(run).index(text) if text is not None else len(run), properties)
    for kind in ("f", "rad", "sSup", "sSub", "sSubSup", "bar", "acc", "d", "borderBox", "eqArr", "m"):
        for element in equation.iter(qn("m:"+kind)):
            properties = element.find(qn("m:"+kind+"Pr"))
            if properties is None:
                properties = OxmlElement("m:"+kind+"Pr")
                element.insert(0, properties)
            control = OxmlElement("m:ctrlPr")
            control.append(font_properties(font_pt, "Cambria Math"))
            properties.append(control)
    return equation


def readable_paragraph(paragraph, *, after=4, font_pt=None, font=None):
    font_pt = FONT_PT if font_pt is None else font_pt
    """Direct formatting avoids inherited tiny line boxes in older Hancom."""
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    paragraph.paragraph_format.line_spacing = Pt(font_pt*1.5)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    properties = paragraph._p.get_or_add_pPr()
    grid = properties.find(qn("w:snapToGrid"))
    if grid is None:
        properties.append(w_element("snapToGrid", val="0"))
    existing = properties.find(qn("w:rPr"))
    if existing is not None:
        properties.remove(existing)
    properties.append(font_properties(font_pt, font))


def add_field(paragraph, instruction):
    field = w_element("fldSimple", instr=instruction)
    run = w_element("r")
    text = w_element("t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def add_page_column_rule(header_paragraph, rule_id=1000000):
    """Repeat a solid black divider through every question and answer page."""
    # Hanword ignores VML header shapes and preserves image aspect ratios.
    # Match the actual tall raster to its printed proportions; a stretched
    # square bitmap becomes a dot in older readers. Uniform black stays sharp.
    pixels = BytesIO()
    Image.new("RGB", (4, round(Mm(259).pt*4)), "black").save(pixels, format="PNG")
    pixels.seek(0)
    inline = header_paragraph.add_run().add_picture(pixels, width=Pt(1), height=Mm(259))._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in {"distT": "0", "distB": "0", "distL": "0", "distR": "0",
                        "simplePos": "0", "relativeHeight": "251659264", "behindDoc": "0",
                        "locked": "1", "layoutInCell": "1", "allowOverlap": "1"}.items():
        anchor.set(name, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0"); simple.set("y", "0")
    anchor.append(simple)
    for axis, offset in (("H", int(Mm(105)-Pt(.5))), ("V", int(Mm(20)))):
        position = OxmlElement("wp:position"+axis)
        position.set("relativeFrom", "page")
        coordinate = OxmlElement("wp:posOffset")
        coordinate.text = str(offset)
        position.append(coordinate)
        anchor.append(position)
    anchor.append(deepcopy(inline.find(qn("wp:extent"))))
    anchor.append(OxmlElement("wp:wrapNone"))
    properties = deepcopy(inline.find(qn("wp:docPr")))
    properties.set("id", str(rule_id))
    properties.set("name", "ExamPageColumnRule")
    properties.set("descr", "2단 중앙 구분선")
    anchor.append(properties)
    for name in ("wp:cNvGraphicFramePr", "a:graphic"):
        node = inline.find(qn(name))
        if node is not None:
            anchor.append(deepcopy(node))
    inline.getparent().replace(inline, anchor)


@lru_cache(maxsize=8)
def load_measure_font(font_pt=None):
    font_pt = FONT_PT if font_pt is None else font_pt
    windir = Path(os.environ.get("WINDIR", "C:/Windows"))
    for file in (windir / {"맑은 고딕":"Fonts/malgun.ttf", "바탕":"Fonts/batang.ttc", "돋움":"Fonts/gulim.ttc"}.get(BODY_FONT,"Fonts/malgun.ttf"), windir / "Fonts/malgun.ttf"):
        try:
            return ImageFont.truetype(str(file), size=round(font_pt*4), index=2 if BODY_FONT == "돋움" and file.name == "gulim.ttc" else 0), 4
        except OSError:
            pass
    return None, 1


def string_width(text, font_pt=None):
    font_pt = FONT_PT if font_pt is None else font_pt
    font, scale = load_measure_font(font_pt)
    if font:
        return font.getlength(text) / scale
    return text_width_em(text) * font_pt


def rich_paragraphs(text):
    """Split prose/newlines and display math into real Word paragraphs."""
    paragraphs, current = [], []
    for kind, content in split_math(text or ""):
        if kind == "display":
            if current:
                paragraphs.append((False, current))
                current = []
            paragraphs.append((True, [(kind, content)]))
        elif kind == "text":
            lines = content.split("\n")
            for index, line in enumerate(lines):
                if line:
                    current.append(("text", line))
                if index < len(lines)-1:
                    if current:
                        paragraphs.append((False, current))
                    current = []
        else:
            current.append((kind, content))
    if current:
        paragraphs.append((False, current))
    return paragraphs or [(False, [("text", "")])]


def estimate_paragraph(display, parts, font_pt=None, after=None):
    font_pt = FONT_PT if font_pt is None else font_pt
    line_pt = font_pt*1.5
    width, line_height, height = 0.0, line_pt, 0.0
    for kind, content in parts:
        if kind == "text":
            for char in content:
                advance = string_width(char, font_pt)
                if width + advance > COL_WIDTH_PT - 3:
                    height += line_height
                    width, line_height = 0.0, line_pt
                width += advance
        else:
            math_width, math_height = dimensions(parse_latex(content))
            # Use the same Cambria Math size as the current body/note context.
            advance = math_width * font_pt
            required_height = max(line_pt, math_height * font_pt * 1.08)
            if advance > COL_WIDTH_PT - 3:
                raise MathSyntaxError("한 단보다 긴 수식입니다. aligned 환경 또는 여러 수식 줄로 나누어 주세요: " + content[:100])
            if width and width + advance > COL_WIDTH_PT - 3:
                height += line_height
                width, line_height = 0.0, line_pt
            width += advance
            line_height = max(line_height, required_height)
    # Native Hancom/Word rendering includes extra font ascent and equation
    # bearings. Keep a margin even when fonts and line pitch are explicit.
    return (height + line_height) * 1.08 + (after if after is not None else 8 if display else 4)


def estimate_text(text, font_pt=None, after=None, wrap_math=False):
    font_pt = FONT_PT if font_pt is None else font_pt
    return sum(estimate_paragraph(display, parts, font_pt, after) for display, parts in output_rich_paragraphs(text, font_pt, wrap_math))


WEB_FIGURE_SIZE_PT = {}


def figure_size(path):
    if not path:
        return None
    image_path = Path(path)
    if not image_path.is_file():
        raise ValueError(f"도형 파일을 찾을 수 없습니다: {image_path.name}")
    with Image.open(image_path) as picture:
        width, height = picture.size
        if width < 1 or height < 1:
            raise ValueError("도형의 크기가 잘못되었습니다.")
    measured = WEB_FIGURE_SIZE_PT.get(str(path))
    if measured is not None:
        try:
            measured_width, measured_height = float(measured["width"]), float(measured["height"])
        except (TypeError, KeyError, ValueError):
            raise ValueError("웹 그림 크기가 잘못되었습니다.")
        if not all(math.isfinite(v) and v > 0 for v in (measured_width, measured_height)):
            raise ValueError("웹 그림 크기가 잘못되었습니다.")
        fit = min(1, (COL_WIDTH_PT - 8) / measured_width)
        return measured_width * fit, measured_height * fit
    max_width = COL_WIDTH_PT - 8
    max_height = 58 * 72 / 25.4
    scale = min(max_width / width, max_height / height)
    return width * scale, height * scale


def choice_rows(choices):
    labels = ["①", "②", "③", "④", "⑤", "⑥", "⑦", "⑧", "⑨", "⑩"]
    rows = []
    for index, choice in enumerate(choices):
        label = labels[index] if index < len(labels) else f"({index+1})"
        # A separate paragraph prevents 5 long formulas from overflowing a line.
        # Recognition may retain printed choice labels. Store the source as-is,
        # but emit exactly one label, matching the preview.
        content = re.sub(r"^\s*[①②③④⑤⑥⑦⑧⑨⑩]\s*", "", str(choice))
        rows.append(f"{label} {content}")
    return rows


def subquestion_parts(value):
    text = value or ""
    masked = re.sub(r"\$\$[\s\S]*?\$\$|(?<!\\)\$[^$\n]*?(?<!\\)\$|\\\[[\s\S]*?\\\]|\\\([\s\S]*?\\\)", lambda m: " " * len(m[0]), text)
    marks = list(re.finditer(r"(?:^|\s)[(（](\d{1,2})[)）](?=\s|$)", masked))
    if len(marks) < 2 or any(int(m[1]) != i+1 for i, m in enumerate(marks)):
        return [text]
    cuts = [m.start() + re.search(r"[(（]", m[0]).start() for m in marks[1:]]
    starts, ends = [0] + cuts, cuts + [len(text)]
    return [text[start:end].rstrip() for start, end in zip(starts, ends)]


def append_choices(document, question, keep=True):
    choices = choice_rows(question.get("choices", []))
    if question.get("webChoiceColumns") != 2 or not choices:
        return [p for text in choices for p in append_rich(document, text, keep=keep)]
    # The web preview uses two independent cells, not a tab-separated line.
    # Keep equations editable and use the same row order and column width.
    table = document.add_table(rows=(len(choices)+1)//2, cols=2)
    table.autofit = False
    for column in table.columns:
        column.width = Pt(COL_WIDTH_PT/2)
    borders = w_element("tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(w_element(side, val="nil"))
    table._tbl.tblPr.append(borders)
    margins = w_element("tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        margins.append(w_element(side, w="0", type="dxa"))
    table._tbl.tblPr.append(margins)
    for i, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(w_element("cantSplit"))
        for j, cell in enumerate(row.cells):
            cell.width = Pt(COL_WIDTH_PT/2)
            original = cell.paragraphs[0]
            index = i*2+j
            paragraphs = append_rich(cell, choices[index], keep=keep and i<len(table.rows)-1) if index<len(choices) else [original]
            if paragraphs != [original]:
                cell._tc.remove(original._p)
            for p in paragraphs:
                readable_paragraph(p, after=6 if i<len(table.rows)-1 else 0)
                p.paragraph_format.keep_with_next = keep and i<len(table.rows)-1
    # A root paragraph terminates the question's keep chain without reserving
    # another text line after the table. It is also the existing border anchor.
    tail = document.add_paragraph()
    readable_paragraph(tail, after=0)
    tail.paragraph_format.line_spacing = Pt(1)
    tail.paragraph_format.keep_with_next = False
    return [tail]


def flow_parts(value):
    text = value or ""
    legacy, cuts, offset = subquestion_parts(text), {}, 0
    for i in range(1, len(legacy)):
        offset = text.index(legacy[i], offset + len(legacy[i-1]))
        cuts[offset] = f"part:{i}"
    masked = re.sub(r"\$\$[\s\S]*?\$\$|(?<!\\)\$[^$\n]*?(?<!\\)\$|\\\[[\s\S]*?\\\]|\\\([\s\S]*?\\\)", lambda m: " " * len(m[0]), text)
    gap = 0
    for match in re.finditer(r"\r?\n[\t ]*\r?\n(?:[\t ]*\r?\n)*[\t ]*|\r?\n[\t ]*(?=[(（]\d{1,2}[)）]\s)", masked):
        at = match.end()
        if at < len(text) and at not in cuts:
            gap += 1
            cuts[at] = f"gap:{gap}"
    line = 0
    for match in re.finditer(r"\r?\n[\t ]*", masked):
        at = match.end()
        if at < len(text) and masked[at] not in "\r\n" and at not in cuts and match.start() not in cuts and (match.start() == 0 or masked[match.start()-1] not in "\r\n"):
            line += 1
            cuts[at] = f"line:{line}"
    starts = [0] + sorted(n for n in cuts if n > 0)
    return [(text[start:starts[i+1] if i+1 < len(starts) else len(text)].rstrip(), cuts[start] if i else "before") for i, start in enumerate(starts)]


def figure_slot(question, identifier, count):
    placements = question.get("figurePlacements")
    slot = placements.get(identifier) if isinstance(placements, dict) else None
    if slot in ("before", "after") or slot in [s for _, s in flow_parts(question.get("body"))]:
        return slot
    return "before" if identifier == "diagram" and question.get("diagramPosition") == "before" else "after"


def box_slot(question):
    slot = question.get("boxSlot", "after")
    return slot if slot in ("before", "after") or slot in [s for _, s in flow_parts(question.get("body"))] else "after"


def estimate_question(question, workspace_lines):
    from structured_docx import active, estimate
    if active(question):
        return LINE_PT + 7 + estimate(question, dict(figure_size=figure_size, estimate_text=estimate_text, string_width=string_width, column_width=COL_WIDTH_PT, line_pt=LINE_PT)) + sum(estimate_text(row) for row in choice_rows(question.get('choices', []))) + workspace_lines*LINE_PT
    parts = [text for text, _ in flow_parts(question["body"])]
    height = LINE_PT + 7 + sum(estimate_text(part) for part in parts) + LINE_PT*(len(parts)-1)
    height += sum(figure_size(p)[1] + 9 for p in question.get("materialPaths", []) if figure_size(p))
    height += sum(estimate_text(row) for row in question.get("statementBox", [])) + (14 if question.get("statementBox") else 0)
    if question.get("bodyBorder"):
        height += 14
    image_size = figure_size(question.get("diagramPath"))
    if image_size:
        height += image_size[1] + 9
    for row in choice_rows(question.get("choices", [])):
        height += estimate_text(row)
    return height + max(0, workspace_lines)*LINE_PT + 9


def validate_snapshot(snapshot):
    if not isinstance(snapshot, dict):
        raise ValueError("잘못된 문서 입력 형식입니다.")
    source = snapshot.get("questions")
    if not isinstance(source, list) or not source:
        raise ValueError("문서에 추가할 문제를 먼저 선택해 주세요.")
    if len(source) > 1000:
        raise ValueError("문서는 한 번에 최대 1000문제까지 출력할 수 있습니다.")
    ids = set()
    questions = []
    for raw in source:
        if not isinstance(raw, dict):
            raise ValueError("문제 형식이 잘못되었습니다.")
        question = dict(raw)
        if question.get("include") is False:
            continue
        if not isinstance(question.get("id"), str) or not question["id"] or question["id"] in ids:
            raise ValueError("문제 ID가 없거나 중복되었습니다.")
        ids.add(question["id"])
        if question.get("kind") not in ("original", "variant"):
            raise ValueError("문제 종류가 잘못되었습니다.")
        if not isinstance(question.get("sourceId"), str) or not question["sourceId"]:
            raise ValueError("문제의 원문 연결 정보가 없습니다.")
        for field in ("body", "answer", "solution"):
            if not isinstance(question.get(field), str):
                raise ValueError(f"{field} 항목은 문자열이어야 합니다.")
            if len(question[field]) > 100000:
                raise ValueError(f"{field} 항목이 너무 깁니다.")
            # Parse all fields before writing a DOCX, including every endnote.
            try:
                for _, parts in rich_paragraphs(question[field]):
                    for kind, content in parts:
                        if kind != "text":
                            parse_latex(content)
            except MathSyntaxError as error:
                caption = {"body": "본문", "answer": "정답", "solution": "상세 풀이"}[field]
                raise MathSyntaxError(f"{len(questions)+1}번 {caption}: {error}") from error
        if not question["body"].strip():
            raise ValueError("본문이 비어 있는 문제는 출력할 수 없습니다.")
        if not isinstance(question.get("choices", []), list) or not all(isinstance(c, str) for c in question.get("choices", [])):
            raise ValueError("선택지 형식이 잘못되었습니다.")
        if len(question.get("choices", [])) > 30:
            raise ValueError("선택지는 최대 30개입니다.")
        for index, choice in enumerate(question.get("choices", []), 1):
            try:
                estimate_text(choice)
            except MathSyntaxError as error:
                raise MathSyntaxError(f"{len(questions)+1}번 선택지 {index}: {error}") from error
        if not isinstance(question.get("statementBox", []), list) or len(question.get("statementBox", [])) > 50 or not all(isinstance(c, str) for c in question.get("statementBox", [])):
            raise ValueError("보기 상자 형식이 잘못되었습니다.")
        if not isinstance(question.get("bodyBorder", False), bool):
            raise ValueError("원본 문제 테두리 형식이 잘못되었습니다.")
        for statement in question.get("statementBox", []):
            estimate_text(statement)
        if question.get("layout", "auto") not in ("auto", "half", "full"):
            raise ValueError("문제 배치 설정이 잘못되었습니다.")
        questions.append(question)
    if not questions:
        raise ValueError("문서에 추가된 문제가 없습니다.")
    return questions


def plan_layout(questions, settings):
    mode = str(settings.get("layout", "auto"))
    if mode not in ("auto", "4", "2"):
        raise ValueError("페이지 배치 설정이 잘못되었습니다.")
    lines = settings.get("workspaceLines", 2)
    if not isinstance(lines, int) or isinstance(lines, bool) or not 0 <= lines <= 20:
        raise ValueError("풀이 공간은 0~20줄로 설정해 주세요.")
    heights = {}
    for number, question in enumerate(questions, 1):
        try:
            heights[question["id"]] = estimate_question(question, lines)
        except MathSyntaxError as error:
            raise MathSyntaxError(f"{number}번 문제: {error}") from error
    if settings.get("measuredPages") is not None:
        pages, used = [], []
        fragments = settings.get('questionFragments', {})
        if not isinstance(fragments, dict) or any(qid not in heights for qid in fragments):
            raise ValueError('분할 문항의 원본 식별자가 잘못되었습니다.')
        expected = []
        from structured_docx import tree as validate_fragment
        for question in questions:
            qid = question['id']
            if qid not in fragments:
                expected.append((qid, None))
                continue
            parts = fragments[qid]
            if not isinstance(parts, list) or len(parts) < 2:
                raise ValueError('분할 문항의 내용이 누락되었습니다.')
            end, total = 0, None
            for index, part in enumerate(parts):
                if not isinstance(part, dict) or any(type(part.get(k)) is not int for k in ('start', 'end', 'total')):
                    raise ValueError('분할 문항의 범위가 잘못되었습니다.')
                total = part['total'] if total is None else total
                if part['total'] != total or part['start'] != end or not end < part['end'] <= total:
                    raise ValueError('분할 문항의 내용이 중복되거나 누락되었습니다.')
                validate_fragment(part['layoutDocument'])
                end = part['end']
                expected.append((qid, index))
            if end != total:
                raise ValueError('분할 문항의 마지막 내용이 누락되었습니다.')
        overflow_pages = settings.get('measuredOverflowPages')
        if overflow_pages is None:
            overflow_pages = [False]*len(settings['measuredPages'])
        if not isinstance(overflow_pages, list) or len(overflow_pages) != len(settings['measuredPages']) or any(type(v) is not bool for v in overflow_pages):
            raise ValueError('측정 페이지의 큰 문항 정보가 잘못되었습니다.')
        for columns in settings["measuredPages"]:
            if not isinstance(columns, list) or len(columns) != 2:
                raise ValueError("측정 페이지는 두 단이어야 합니다.")
            if not all(isinstance(column, list) for column in columns):
                raise ValueError("측정 단의 문항 목록이 잘못되었습니다.")
            if settings.get("quadrantLayout") and any(len(column) > 2 for column in columns):
                raise ValueError("기본 시험지 배치는 한 단에 최대 2문항입니다. 미리보기를 새로고침해 주세요.")
            positions = []
            for col, ids in enumerate(columns):
                for slot, ref in enumerate(ids):
                    if isinstance(ref, str):
                        position = {'questionId': ref, 'column': col, 'slot': slot}
                    elif isinstance(ref, dict) and isinstance(ref.get('questionId'), str) and type(ref.get('fragmentIndex')) is int:
                        position = dict(questionId=ref['questionId'], fragmentIndex=ref['fragmentIndex'], column=col, slot=slot)
                    else:
                        raise ValueError('측정 배치의 문항 참조가 잘못되었습니다.')
                    positions.append(position)
            used.extend((p['questionId'], p.get('fragmentIndex')) for p in positions)
            overflow = overflow_pages[len(pages)]
            pages.append({"positions": positions, "count": len(positions), "capacity": 4 if settings.get("quadrantLayout") else 2,
                          "questionIds": [p["questionId"] for p in positions], "overflow": overflow})
        if used != expected:
            raise ValueError("측정 배치와 문항 순서가 다릅니다.")
        return pages, heights, ["브라우저 측정 배치 사용. 편집 프로그램의 글꼴·수식 폭에 따라 최종 줄바꿈은 달라질 수 있습니다."]
    pages, warnings = [], []
    index = 0
    while index < len(questions):
        current = questions[index]
        if heights[current["id"]] > COL_HEIGHT_PT - 20:
            estimated_columns = max(2, math.ceil(heights[current["id"]] / (COL_HEIGHT_PT-20)))
            pages.append({"count": 1, "capacity": 2, "questionIds": [current["id"]],
                          "overflow": True, "estimatedColumns": estimated_columns,
                          "positions": [{"questionId": current["id"], "column": 0, "slot": 0}]})
            warnings.append(f"{index+1}번은 한 단보다 길어 다음 단/페이지로 이어집니다. 출력에서 페이지 나눔을 확인해 주세요.")
            index += 1
            continue
        # A column owns its midpoint slots, independently of the neighbouring
        # column and of writing-space settings. Browser measurements take
        # precedence; this path preserves the same rule for direct exports.
        positions = []
        column, slot = 0, 0
        slots = 1 if mode == "2" else 2
        for q in questions[index:]:
            height = heights[q["id"]]
            if height > COL_HEIGHT_PT - 20:
                break
            full = q.get("layout") == "full" or height > COL_HEIGHT_PT/slots-13.5
            if slot >= slots or (full and slot):
                column, slot = column+1, 0
            if column >= 2:
                break
            positions.append({"questionId": q["id"], "column": column, "slot": slot})
            slot = slots if full else slot+1
        pages.append({"count": len(positions), "capacity": slots*2,
                      "questionIds": [p["questionId"] for p in positions],
                      "overflow": False, "positions": positions})
        index += len(positions)
    return pages, heights, warnings


def configure_document(document, title, show_student_name_line=False):
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(20), Mm(18)
    section.left_margin = section.right_margin = Mm(15)
    section.header_distance, section.footer_distance = Mm(8), Mm(8)
    columns = section._sectPr.find(qn("w:cols"))
    columns.set(qn("w:num"), "2")
    columns.set(qn("w:space"), str(round(9*1440/25.4)))
    columns.set(qn("w:sep"), "1")
    # Asian document grids can enlarge a text line to the next grid line in
    # Hancom. Disable the grid explicitly rather than estimating that expansion.
    grid = section._sectPr.find(qn("w:docGrid"))
    if grid is not None:
        section._sectPr.remove(grid)
    # python-docx's bundled template carries themed Title/Subtitle borders.
    # Remove them so a printed math sheet has no imported blue header rule.
    for border in list(document.styles.element.xpath(".//w:pBdr")):
        border.getparent().remove(border)
    normal = document.styles["Normal"]
    set_font(normal)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    normal.paragraph_format.line_spacing = Pt(LINE_PT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True
    normal.element.get_or_add_pPr().append(w_element("snapToGrid", val="0"))
    for name, size in (("Title", 13), ("Header", 10), ("Footer", 9)):
        set_font(document.styles[name], size)
    for name, size in (("Question Label", FONT_PT), ("Endnote Text", NOTE_FONT_PT), ("Endnote Heading", NOTE_FONT_PT)):
        if name not in document.styles:
            document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = document.styles[name]
        style.base_style = normal
        set_font(style, size, NOTE_FONT if name.startswith("Endnote") else BODY_FONT)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        style.paragraph_format.line_spacing = Pt(size*1.5)
    # The label shares the first prose paragraph; only its own run is bold.
    document.styles["Question Label"].font.bold = False
    document.styles["Question Label"].paragraph_format.space_after = Pt(7)
    document.styles["Question Label"].paragraph_format.keep_with_next = True
    document.styles["Endnote Heading"].font.bold = True
    document.styles["Endnote Heading"].paragraph_format.keep_with_next = True
    if "Endnote Reference" not in document.styles:
        document.styles.add_style("Endnote Reference", WD_STYLE_TYPE.CHARACTER)
    note_style = document.styles["Endnote Reference"]
    set_font(note_style, 9)
    note_style.font.superscript = True
    header = section.header.paragraphs[0]
    header.style = document.styles["Title"]
    header.text = title.strip() or "수학 유사문제"
    header.paragraph_format.space_after = Pt(0)
    add_page_column_rule(header)
    if show_student_name_line:
        section.different_first_page_header_footer = True
        first_header = section.first_page_header.paragraphs[0]
        first_header.style = document.styles["Title"]
        first_header.text = header.text
        first_header.paragraph_format.space_after = Pt(0)
        first_header.paragraph_format.tab_stops.add_tab_stop(Mm(180), WD_TAB_ALIGNMENT.RIGHT)
        set_font(first_header.add_run("\t이름: ________________"), 10, BODY_FONT)
        add_page_column_rule(first_header, rule_id=1000001)
        first_footer = section.first_page_footer.paragraphs[0]
        first_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(first_footer, "PAGE")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE")
    document.core_properties.title = title.strip() or "수학 유사문제"
    document.core_properties.subject = "원본문제와 유사문제 및 상세 풀이"
    document.core_properties.author = ""
    settings = document.settings.element
    # Prevent final-page column balancing from squeezing writing space.
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = w_element("compat")
        settings.append(compat)
    compat.append(w_element("doNotBalanceTextColumns"))


def solution_alignment_rows(source):
    """Read top-level alignment rows without cutting a group or nested array."""
    match = re.fullmatch(r"\\begin\{(aligned|gathered)\}([\s\S]*)\\end\{\1\}", source.strip())
    if not match:
        return None
    body, rows, cells, start, depth, environments, i = match[2], [], [], 0, 0, 0, 0
    while i < len(body):
        command = re.match(r"\\[A-Za-z]+|\\.", body[i:]) if body[i] == "\\" else None
        token = command.group() if command else body[i]
        if token == r"\begin":
            environments += 1
        elif token == r"\end":
            environments -= 1
        if depth == 0 and environments == 0 and token in (r"\\", "&"):
            cells.append(body[start:i].strip())
            start = i + len(token)
            if token == r"\\":
                rows.append(cells)
                cells = []
        elif not command:
            if token == "{": depth += 1
            elif token == "}": depth -= 1
        i += len(token)
    cells.append(body[start:].strip())
    rows.append(cells)
    # Flatten one equation per row only. Multiple independent alignment pairs
    # must not be concatenated into a different mathematical expression.
    if any(len(row) > 2 or (len(row) == 2 and row[0] and not re.match(r"^(?:[=<>+\-]|\\(?:leq?|geq?|neq?|approx|equiv|sim)\b)", row[1])) for row in rows):
        return None
    return [''.join(row) for row in rows if any(row)]


def solution_math_lines(source, font_pt):
    """Split only top-level binary relations; keep groups/fractions intact.

    LibreOffice's native-math advances can be much wider than our planning
    estimate. Reserve that measured headroom instead of shrinking or clipping.
    This is opt-in for web solution copies, never a change to saved LaTeX.
    """
    limit = COL_WIDTH_PT * .45
    width = lambda value: dimensions(parse_latex(value))[0] * font_pt
    if width(source) <= limit:
        return [source]
    tree = parse_latex(source)
    if len(tree.children) == 1 and tree.children[0].kind == "box" and width(source) <= COL_WIDTH_PT * .65:
        # A single answer box is indivisible. Retain its border and reserve
        # more than 35% of the column for native glyph and border bearings.
        return [source]
    alignment = solution_alignment_rows(source)
    if alignment is not None:
        return [part for row in alignment for part in solution_math_lines(row, font_pt)]
    cuts, braces, brackets, parens, left_right, environments = [], 0, 0, 0, 0, 0
    i = 0
    while i < len(source):
        char = source[i]
        if char == "\\":
            match = re.match(r"\\[A-Za-z]+|\\.", source[i:])
            token = match.group() if match else char
            if token == r"\left":
                left_right += 1
            elif token == r"\right":
                left_right = max(0, left_right-1)
            elif token == r"\begin":
                environments += 1
            elif token == r"\end":
                environments -= 1
            elif token in (r"\times", r"\cdot", r"\le", r"\leq", r"\ge", r"\geq", r"\ne", r"\neq") and not (braces or brackets or parens or left_right or environments) and i:
                cuts.append(i)
            i += len(token)
            continue
        if char == "{": braces += 1
        elif char == "}": braces -= 1
        elif char == "[": brackets += 1
        elif char == "]": brackets -= 1
        elif char == "(": parens += 1
        elif char == ")": parens -= 1
        elif char in "=+-" and not (braces or brackets or parens or left_right or environments) and i and source[i-1] not in "=+-^_":
            cuts.append(i)
        i += 1
    rows, start = [], 0
    for end in [*cuts, len(source)]:
        if end <= start:
            continue
        candidate = source[start:end]
        # A continuation keeps its operator, so equality/operation is explicit.
        if width(candidate) <= limit:
            continue
        previous = next((cut for cut in reversed(cuts) if start < cut < end and width(source[start:cut]) <= limit), None)
        if previous is None:
            # An indivisible object is never silently cropped or reduced.
            if width(candidate) > limit:
                raise MathSyntaxError("상세 풀이의 수식을 한 단 안에서 나눌 수 없습니다. 수식 줄을 직접 나누어 주세요: " + candidate[:100])
            continue
        rows.append(source[start:previous].strip())
        start = previous
    rows.append(source[start:].strip())
    if any(width(row) > limit for row in rows):
        raise MathSyntaxError("상세 풀이의 수식이 한 단보다 큽니다. 수식 줄을 나누어 주세요: " + source[:100])
    return rows


def output_rich_paragraphs(text, font_pt, wrap_math=False):
    rich_parts = []
    for display, parts in rich_paragraphs(text):
        if wrap_math:
            current = []
            for kind, content in parts:
                rows = solution_math_lines(content, font_pt) if kind != "text" else [content]
                if len(rows) > 1:
                    if current:
                        rich_parts.append((display, current)); current = []
                    operators = {r"\times":"×", r"\cdot":"·", r"\le":"≤", r"\leq":"≤", r"\ge":"≥", r"\geq":"≥", r"\ne":"≠", r"\neq":"≠"}
                    for row in rows:
                        # Match whole TeX commands: \left and \leftarrow are
                        # not the relation \le followed by ordinary letters.
                        match = re.match(r"^(=|\+|-|\\(?:times|cdot|leq?|geq?|neq?)(?![A-Za-z]))", row)
                        if match:
                            # Writer rejects standalone OMML that starts with a
                            # binary operator. The continuation operator remains
                            # ordinary editable text alongside native RHS math.
                            operator = match.group()
                            rich_parts.append((True, [("text", operators.get(operator, operator)+" "), ("inline", row[len(operator):])]))
                        else:
                            rich_parts.append((True, [("inline", row)]))
                else:
                    current.append((kind, content))
            if current:
                rich_parts.append((display, current))
        else:
            rich_parts.append((display, parts))
    return rich_parts


def append_rich(document, text, style=None, keep=True, font_pt=None, after=None, font=None, wrap_math=False):
    font_pt = FONT_PT if font_pt is None else font_pt
    paragraphs = []
    for display, parts in output_rich_paragraphs(text, font_pt, wrap_math):
        paragraph = document.add_paragraph(style=style)
        readable_paragraph(paragraph, after=after if after is not None else 8 if display else 4, font_pt=font_pt, font=font)
        paragraph.paragraph_format.keep_with_next = keep
        paragraph.paragraph_format.keep_together = keep
        if display:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for kind, content in parts:
            if kind == "text":
                set_font(paragraph.add_run(content), font_pt, font)
            else:
                paragraph._p.append(formatted_math(content, display=display, font_pt=font_pt))
        paragraphs.append(paragraph)
    return paragraphs


def add_solution_reference(paragraph, number):
    link = w_element("hyperlink", anchor=f"solution_{number}", history="1")
    run = w_element("r")
    properties = font_properties(1)
    properties.append(w_element("vanish"))
    properties.append(w_element("color", val="FFFFFF"))
    run.append(properties)
    text = w_element("t")
    text.text = str(number)
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def inline_question_label(heading, content, question):
    """Join only a leading prose paragraph; retain all intentional blocks."""
    if not content or question.get("bodyBorder"):
        return heading
    first = content[0]
    if (heading._p.getnext() is not first._p or
            first._p.find('.//' + qn('w:drawing')) is not None or
            first._p.find('.//' + qn('m:oMathPara')) is not None or
            first._p.find('w:pPr/w:pBdr', first._p.nsmap) is not None):
        return heading
    first.style = 'Question Label'
    first.paragraph_format.page_break_before = heading.paragraph_format.page_break_before
    prefix = [child for child in heading._p if child.tag != qn('w:pPr')]
    gap = w_element('r'); text = w_element('t'); text.text = ' '; text.set(qn('xml:space'), 'preserve'); gap.append(text)
    for child in reversed(prefix + [gap]):
        first._p.insert(1 if first._p.pPr is not None else 0, child)
    heading._p.getparent().remove(heading._p)
    return first


def finish_question_content(heading, content, question, number, student):
    if question.get('_continuation') and content:
        first = content[0]
        first.paragraph_format.page_break_before = heading.paragraph_format.page_break_before
        for child in list(heading._p):
            if child.tag != qn('w:pPr'):
                first._p.insert(1, child)
        heading._p.getparent().remove(heading._p)
        heading = first
    else:
        heading = inline_question_label(heading, content, question)
    if not student and question.get('_fragmentLast', True):
        # content follows actual document order, including nested choices and
        # figures. Append to its final paragraph, before any writing space.
        add_solution_reference(content[-1] if content else heading, number)
    return heading


def create_solutions(document, questions, wrap_math=False):
    """Flow through the answer section's real columns in every DOCX reader."""
    for number, question in enumerate(questions, 1):
        heading = document.add_paragraph(style="Endnote Heading")
        readable_paragraph(heading, after=7, font_pt=NOTE_FONT_PT, font=NOTE_FONT)
        heading._p.append(w_element("bookmarkStart", id=number, name=f"solution_{number}"))
        printed = question.get("printedNumber") or str(number)
        set_font(heading.add_run(f"{printed}번 정답  "), NOTE_FONT_PT, NOTE_FONT)
        for kind, content in split_math(question["answer"] or "정답·풀이 미생성"):
            if kind == "text":
                set_font(heading.add_run(content), NOTE_FONT_PT, NOTE_FONT)
            else:
                heading._p.append(formatted_math(content, font_pt=NOTE_FONT_PT))
        heading._p.append(w_element("bookmarkEnd", id=number))
        heading.paragraph_format.keep_with_next = True
        # Let detailed solutions continue naturally in the next column/page.
        # Keep only individual paragraphs together, never the entire solution.
        from solution_guide_docx import render as render_solution_guide
        guide_question = dict(question, _guideWrapMath=wrap_math)
        paragraphs = render_solution_guide(document, guide_question, append_rich, set_font, NOTE_FONT_PT, NOTE_FONT, 0)
        if paragraphs is None:
            paragraphs = append_rich(document, question["solution"] or "정답·풀이 미생성",
                                     style="Endnote Text", keep=False, font_pt=NOTE_FONT_PT, after=0, font=NOTE_FONT, wrap_math=wrap_math)
        for paragraph in paragraphs:
            paragraph.paragraph_format.keep_together = False
            paragraph.paragraph_format.widow_control = True
        if paragraphs:
            # One blank line between solutions, without a whole-solution keep chain.
            paragraphs[-1].paragraph_format.space_after = Pt(NOTE_LINE_PT if number < len(questions) else 0)


def add_question(document, question, number, original_numbers, natural_height, target_height, workspace_lines, overflow=False, new_page=False, new_column=False, student=False, show_labels=True, separate_column_break=False, terminal_workspace=False):
    if question["kind"] == "original":
        label = "원본문제"
    else:
        original_number = original_numbers.get(question["sourceId"])
        label = f"유사문제 · 원문 {original_number}번" if original_number else "유사문제 · 원문 별도"
    if new_column and separate_column_break:
        # Writer can discard an inline column break while balancing the final
        # section. A minimal separate paragraph preserves measured web columns.
        column_marker = document.add_paragraph()
        readable_paragraph(column_marker, after=0)
        column_marker.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        column_marker.paragraph_format.line_spacing = Pt(1)
        column_marker.paragraph_format.keep_with_next = True
        column_marker.add_run().add_break(WD_BREAK.COLUMN)
    heading = document.add_paragraph(style="Question Label")
    readable_paragraph(heading, after=7)
    heading.paragraph_format.page_break_before = new_page
    if new_column and not separate_column_break:
        # A separate break paragraph leaves an empty line at the top
        # of the right column in Hancom. Keep the break and label in one line.
        heading.add_run().add_break(WD_BREAK.COLUMN)
    printed = str(question.get("printedNumber") or number)
    printed_label = printed + ("." if printed.isdigit() else "")
    if question.get("originalPoints") is not None and not question.get("pointsAtBodyEnd"):
        original_points = str(question["originalPoints"]).strip()
        printed_label += f" ({original_points if '점' in original_points else original_points + '점'})"
    label_run = heading.add_run('' if question.get('_continuation') else (f"{printed_label} [{label}]" if show_labels else printed_label))
    set_font(label_run)
    label_run.bold = True
    from structured_docx import active, render
    if active(question):
        api = dict(w_element=w_element, font_pt=FONT_PT, readable_paragraph=readable_paragraph,
                   formatted_math=formatted_math, figure_size=figure_size, string_width=string_width, set_font=set_font)
        content = render(document, question, api, COL_WIDTH_PT, keep=not overflow)
        # The proof box may span a column, but its short option list is still
        # one group. Move all options together instead of splitting 1/2 and 3/4/5.
        content.extend(append_choices(document, question, keep=True))
        heading = finish_question_content(heading, content, question, number, student)
        if content:
            content[-1].paragraph_format.keep_with_next = False
        if question.get("quadrantFill") and not overflow:
            mm = question.get("workspaceMm", 0)
            if not isinstance(mm, (int, float)) or not math.isfinite(mm) or not 0 <= mm <= 200:
                raise ValueError("풀이 공간은 0~200mm입니다.")
            padding = max(mm*72/25.4, target_height-natural_height)
            if padding > 0:
                writing = document.add_paragraph()
                readable_paragraph(writing, after=max(0, padding-LINE_PT))
                writing.paragraph_format.line_spacing = Pt(min(LINE_PT, padding))
                writing.paragraph_format.keep_with_next = False
                writing.paragraph_format.keep_together = False
                return writing
        if "workspaceMm" in question and not overflow:
            mm = question["workspaceMm"]
            if not isinstance(mm, (int, float)) or not math.isfinite(mm) or not 0 <= mm <= 200:
                raise ValueError("풀이 공간은 0~200mm입니다.")
            if mm:
                if terminal_workspace and content:
                    last = content[-1]
                    last.paragraph_format.space_after = Pt((last.paragraph_format.space_after.pt if last.paragraph_format.space_after else 0) + mm*72/25.4)
                    last.paragraph_format.keep_with_next = False
                    return last
                writing = document.add_paragraph()
                readable_paragraph(writing, after=max(0, mm*72/25.4-LINE_PT))
                writing.paragraph_format.line_spacing = Pt(min(LINE_PT, mm*72/25.4))
                writing.paragraph_format.keep_with_next = False
                return writing
        elif workspace_lines and not overflow:
            writing = document.add_paragraph()
            readable_paragraph(writing, after=max(0, workspace_lines*LINE_PT-LINE_PT))
            writing.paragraph_format.keep_with_next = False
            return writing
        return content[-1] if content else heading
    content, part_starts, figure_paragraphs = [], [], []
    flow = flow_parts(question["body"])
    for index, (part, _) in enumerate(flow):
        paragraphs = append_rich(document, part, keep=not overflow)
        if index:
            paragraphs[0].paragraph_format.space_before = Pt(LINE_PT)
        part_starts.append(paragraphs[0])
        content.extend(paragraphs)
    if question.get("pointsAtBodyEnd") and question.get("originalPoints") is not None and content:
        original_points = str(question["originalPoints"]).strip()
        if original_points:
            suffix = original_points if "점" in original_points else original_points + "점"
            set_font(content[-1].add_run(f" ({suffix})"))
    box_paragraphs = []
    if question.get("statementBox"):
        box_paragraphs = append_rich(document, "\n".join(question["statementBox"]), keep=not overflow)
        for index, paragraph in enumerate(box_paragraphs):
            borders = w_element("pBdr")
            sides = ["left", "right"]
            if index == 0:
                sides.append("top")
            if index == len(box_paragraphs) - 1:
                sides.append("bottom")
            for side in sides:
                borders.append(w_element(side, val="single", sz="6", space="6", color="7F8C7A"))
            paragraph._p.get_or_add_pPr().append(borders)
            paragraph.paragraph_format.keep_with_next = index < len(box_paragraphs) - 1
        content.extend(box_paragraphs)
    for material_index, material_path in enumerate(question.get("materialPaths", [])):
        width, height = figure_size(material_path)
        paragraph = document.add_paragraph()
        readable_paragraph(paragraph)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = not overflow
        paragraph.add_run().add_picture(material_path, width=Pt(width), height=Pt(height))
        drawing = paragraph._p.find(".//" + qn("wp:docPr"))
        if drawing is not None:
            drawing.set("descr", f"{number}번 문제의 공통 자료 {material_index+1}")
        content.append(paragraph)
        identifiers = question.get("materialIds", [])
        identifier = identifiers[material_index] if material_index < len(identifiers) else f"material-{material_index}"
        figure_paragraphs.append((identifier, paragraph))
    if question.get("diagramPath"):
        width, height = figure_size(question["diagramPath"])
        paragraph = document.add_paragraph()
        readable_paragraph(paragraph)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = not overflow
        paragraph.add_run().add_picture(question["diagramPath"], width=Pt(width), height=Pt(height))
        drawing = paragraph._p.find(".//" + qn("wp:docPr"))
        if drawing is not None:
            drawing.set("descr", f"{number}번 문제의 도형")
        content.append(paragraph)
        figure_paragraphs.append(("diagram", paragraph))
    if box_paragraphs and box_slot(question) != "after":
        slot = box_slot(question)
        anchor = part_starts[next(i for i, (_, place) in enumerate(flow) if place == slot)]
        for paragraph in box_paragraphs:
            anchor._p.addprevious(paragraph._p)
    for identifier, paragraph in figure_paragraphs:
        slot = figure_slot(question, identifier, len(part_starts))
        if slot != "after":
            anchor = part_starts[next(i for i, (_, place) in enumerate(flow) if place == slot)]
            anchor._p.addprevious(paragraph._p)
    content.extend(append_choices(document, question, keep=not overflow))
    order = {node: index for index, node in enumerate(document._element.body)}
    content.sort(key=lambda paragraph: order[paragraph._p])
    if question.get("bodyBorder") and content:
        for index, paragraph in enumerate(content):
            borders = w_element("pBdr")
            sides = ["left", "right"]
            if index == 0:
                sides.append("top")
            if index == len(content) - 1:
                sides.append("bottom")
            for side in sides:
                borders.append(w_element(side, val="single", sz="6", space="6", color="777777"))
            paragraph._p.get_or_add_pPr().append(borders)
    if question.get("sourceCaption"):
        caption = document.add_paragraph()
        readable_paragraph(caption, after=4, font_pt=8)
        set_font(caption.add_run(str(question["sourceCaption"])), 8, NOTE_FONT)
        content.append(caption)
    heading = finish_question_content(heading, content, question, number, student)
    # End keep-with-next chain here so a whole sheet is never pushed together.
    if content:
        content[-1].paragraph_format.keep_with_next = False
    # Natural height already includes minimum writing room; fill a short top
    # block to half a column, while never shrinking text or an equation to fit.
    desired = max(workspace_lines*LINE_PT, target_height-natural_height+workspace_lines*LINE_PT)
    # Never reserve blank writing space beyond a column. In a flowing, very
    # long question its final position is unknown, so reserve no forced padding.
    content_height = max(0, natural_height-workspace_lines*LINE_PT)
    remaining = max(0, COL_HEIGHT_PT-content_height-24)
    padding = 0 if overflow or workspace_lines == 0 else min(desired, remaining)
    if "workspaceMm" in question:
        mm = question["workspaceMm"]
        if not isinstance(mm, (int, float)) or not math.isfinite(mm) or not 0 <= mm <= 200:
            raise ValueError("풀이 공간은 0~200mm입니다.")
        padding = mm * 72 / 25.4
        if question.get("quadrantFill"):
            padding = max(padding, min(remaining, target_height-content_height))
    if padding <= 0 or (padding < LINE_PT and "workspaceMm" not in question):
        return content[-1] if content else heading
    if terminal_workspace:
        # At a measured column/page boundary, a separate empty writing line can
        # spill onto the next page before the following explicit page break.
        # Paragraph spacing retains the requested writing space without adding
        # an otherwise content-free page to Writer's layout.
        last = content[-1] if content else heading
        last.paragraph_format.space_after = Pt((last.paragraph_format.space_after.pt if last.paragraph_format.space_after else 0) + padding)
        last.paragraph_format.keep_with_next = False
        return last
    writing = document.add_paragraph()
    readable_paragraph(writing, after=max(0, padding-LINE_PT))
    if "workspaceMm" in question and padding < LINE_PT:
        writing.paragraph_format.line_spacing = Pt(padding)
    writing.paragraph_format.keep_with_next = False
    writing.paragraph_format.keep_together = False
    return writing


def export_document(snapshot, output, math_validator=None):
    from export_integrity import verify as verify_export_engine
    verify_export_engine()
    global BODY_FONT, FONT_PT, LINE_PT, NOTE_FONT, NOTE_FONT_PT, NOTE_LINE_PT, WEB_FIGURE_SIZE_PT, COL_WIDTH_PT, COL_HEIGHT_PT, HALF_HEIGHT_PT
    settings = snapshot.get("settings", {})
    if not isinstance(settings, dict):
        raise ValueError("문서 설정 형식이 잘못되었습니다.")
    settings = {**settings, "quadrantLayout": settings.get("quadrantLayout", str(settings.get("layout", "auto")) != "2")}
    import paper_form
    form_pages = paper_form.pages(snapshot)
    COL_WIDTH_PT = ((210-form_pages[0]['leftMm']-form_pages[0]['rightMm']-form_pages[0]['gapMm'])/2 if form_pages else 85.5)*72/25.4
    COL_HEIGHT_PT = settings.get('measuredColumnHeightPt',259*72/25.4)
    if not isinstance(COL_HEIGHT_PT,(int,float)) or isinstance(COL_HEIGHT_PT,bool) or not math.isfinite(COL_HEIGHT_PT) or not 50 < COL_HEIGHT_PT <= 297*72/25.4:
        raise ValueError('측정한 문제 영역 높이가 잘못되었습니다.')
    HALF_HEIGHT_PT = (COL_HEIGHT_PT-18)/2
    guide_mode = settings.get("solutionGuideMode", "figures")
    if guide_mode not in ("figures", "text"):
        raise ValueError("해설 보조그림 출력 설정이 잘못되었습니다.")
    if guide_mode == "text":
        # The browser supplies the complete solution text, but does not yet
        # render desktop-only solution-guide figures. Never mutate saved data.
        snapshot = deepcopy(snapshot)
        for question in snapshot.get("questions", []):
            question.pop("solutionGuide", None)
    WEB_FIGURE_SIZE_PT = settings.get("figureSizePt") or {}
    if not isinstance(WEB_FIGURE_SIZE_PT, dict):
        raise ValueError("웹 그림 크기 설정이 잘못되었습니다.")
    BODY_FONT = settings.get("bodyFont", "맑은 고딕")
    if BODY_FONT not in ("맑은 고딕", "바탕", "돋움"):
        BODY_FONT = "맑은 고딕"
    try:
        FONT_PT = float(settings.get("bodyFontSize", 12))
    except (TypeError, ValueError):
        FONT_PT = 12
    # New web exam sizes use 7..15pt. Existing saved 16..18pt remain readable.
    if not math.isfinite(FONT_PT) or not 7 <= FONT_PT <= 18:
        FONT_PT = 12
    NOTE_FONT = settings.get("solutionFont", BODY_FONT)
    if NOTE_FONT not in ("맑은 고딕", "바탕", "돋움"):
        NOTE_FONT = BODY_FONT
    try:
        NOTE_FONT_PT = float(settings.get("solutionFontSize", 9))
    except (TypeError, ValueError):
        NOTE_FONT_PT = 9
    if not math.isfinite(NOTE_FONT_PT) or not 9 <= NOTE_FONT_PT <= 18:
        NOTE_FONT_PT = 9
    NOTE_LINE_PT = NOTE_FONT_PT * 1.5
    LINE_PT = FONT_PT * 1.5
    load_measure_font.cache_clear()
    student = False  # Answer pages are always included; choose pages when printing.
    if student:
        snapshot = deepcopy(snapshot)
        for question in snapshot.get("questions", []):
            question["answer"] = ""
            question["solution"] = ""
    from export_preflight import check as check_export_formulas
    check_export_formulas(snapshot, rich_paragraphs, solution_math_lines, NOTE_FONT_PT, math_validator)
    questions = validate_snapshot(snapshot)
    from solution_guide_docx import steps as preflight_solution_guide
    for question in questions: preflight_solution_guide(question)
    settings = snapshot.get("settings", {})
    if not isinstance(settings, dict):
        raise ValueError("문서 설정 형식이 잘못되었습니다.")
    pages, heights, warnings = plan_layout(questions, settings)
    # Preflight notes as well as questions before replacing a prior export.
    solution_height = 0
    for number, question in enumerate(questions, 1):
        try:
            solution_height += NOTE_LINE_PT + 7 + estimate_text(question["solution"], NOTE_FONT_PT, after=0, wrap_math=settings.get("wrapSolutionMath", False)) + estimate_text(question["answer"], NOTE_FONT_PT) + (NOTE_LINE_PT if number < len(questions) else 0)
            from solution_guide_docx import proof_lines, compact_step
            for step in preflight_solution_guide(question) or []:
                proofs=proof_lines(question,step)
                solution_height+=sum(estimate_text(line,NOTE_FONT_PT,after=0) for line in proofs)
                if step.get('view'):
                    compact=compact_step(question,step)
                    with Image.open(question['solutionGuideFigurePaths'][step['id']]) as image:
                        picture_height=(128 if compact else 220)*image.height/image.width
                    # An estimate only; Word/Hancom remain the pagination engines.
                    solution_height+=(max(0,picture_height-estimate_text(step['text'],NOTE_FONT_PT,after=0)) if compact else picture_height)+NOTE_LINE_PT
        except MathSyntaxError as error:
            raise MathSyntaxError(f"{number}번 정답 또는 상세 풀이: {error}") from error
    if not student and any(not q["answer"].strip() or not q["solution"].strip() for q in questions):
        warnings.append("정답 또는 상세 풀이가 비어 있는 문제가 있습니다.")
    originals = {q["sourceId"]: index for index, q in enumerate(questions, 1) if q["kind"] == "original"}
    if any(q["kind"] == "variant" and q["sourceId"] not in originals for q in questions):
        warnings.append("출력에서 제외된 원문의 유사문제에는 '원문 별도'를 표시했습니다.")
    warnings.append("페이지 배치는 글자·수식·도형 크기로 계산한 추정값입니다. Word 또는 한글에서 최종 페이지 나눔과 답지를 확인해 주세요.")
    document = Document()
    configure_document(document, str(snapshot.get("title", "수학 유사문제")), settings.get("showStudentNameLine", False))
    by_id = {q["id"]: q for q in questions}
    numbers = {q["id"]: index for index, q in enumerate(questions, 1)}
    lines = settings.get("workspaceLines", 2)
    for page_index, page in enumerate(pages):
        if form_pages:
            section = document.sections[0] if page_index == 0 else document.add_section(WD_SECTION.NEW_PAGE)
            paper_form.configure(section, form_pages[page_index])
            if page_index:
                boundary=document.paragraphs[-1]
                readable_paragraph(boundary,after=0)
                boundary.paragraph_format.line_spacing=Pt(1)
            # Use this page's usable area, not the basic form's 259 mm.
            # The first-page instructions affect only the left column.
            COL_HEIGHT_PT = (297-form_pages[page_index]['topMm']-form_pages[page_index]['bottomMm'])*72/25.4
            HALF_HEIGHT_PT = (COL_HEIGHT_PT-18)/2
        previous_column = 0
        form_groups = [[], []]
        for position_index, position in enumerate(page["positions"]):
            new_column = position["column"] != previous_column
            previous_column = position["column"]
            question = by_id[position["questionId"]]
            if 'fragmentIndex' in position:
                parts = settings['questionFragments'][question['id']]
                index = position['fragmentIndex']
                question = dict(question, layoutDocument=parts[index]['layoutDocument'], layoutMode='structure',
                                choices=[], bodyBorder=False, originalPoints=None, sourceCaption=None,
                                _outputFragment=True, _continuation=index > 0, _fragmentLast=index == len(parts)-1)
            has_next_in_column = page["capacity"] == 4 and position["slot"] == 0 and any(p["column"] == position["column"] and p["slot"] == 1 for p in page["positions"])
            target = (COL_HEIGHT_PT/2 if settings.get("quadrantLayout") else HALF_HEIGHT_PT) if has_next_in_column and (lines > 0 or settings.get("quadrantLayout")) else heights[question["id"]]
            if form_pages and has_next_in_column and settings.get("quadrantLayout"):
                intro = form_pages[page_index]['introMm']*72/25.4 if position['column'] == 0 else 0
                target = (COL_HEIGHT_PT-intro)/2
            placed_question = {**question, "quadrantFill": True, "workspaceMm": question.get('workspaceMm',0)} if not form_pages and has_next_in_column and settings.get("quadrantLayout") else question
            before_nodes=set(document._element.body)
            add_question(document, placed_question, numbers[question["id"]], originals,
                         heights[question["id"]], target, lines, page["overflow"],
                         new_page=not form_pages and page_index > 0 and position_index == 0,
                         new_column=new_column and not form_pages, student=student, show_labels=settings.get("showQuestionLabels", True),
                         separate_column_break=settings.get("measuredPages") is not None,
                         terminal_workspace=settings.get("measuredPages") is not None and not settings.get("quadrantLayout") and
                         (position_index == len(page["positions"])-1 or page["positions"][position_index+1]["column"] != position["column"]))
            if form_pages:
                nodes=[n for n in document._element.body if n not in before_nodes]
                for n in nodes:document._element.body.remove(n)
                form_groups[position['column']].append(nodes)
        if form_pages:
            paper_form.question_page(document,form_groups,COL_HEIGHT_PT,form_pages[page_index],COL_WIDTH_PT)
    # A fresh two-column body section keeps quick answers and detailed solutions together.
    if not student:
        appendix = document.add_section(WD_SECTION.NEW_PAGE)
        if form_pages:
            paper_form.configure(appendix, title=str(snapshot.get("title", "")), title_font_pt=settings.get('nativeHeaderFontPt',10), title_font=BODY_FONT)
            COL_WIDTH_PT = 85.5*72/25.4
        # The first-page name field belongs to the document, not each section.
        appendix.different_first_page_header_footer = False
        for paragraph in document.paragraphs[-1:]:
            readable_paragraph(paragraph, after=0)
            if settings.get("measuredPages") is not None:
                paragraph.paragraph_format.line_spacing = Pt(1)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.keep_with_next = False
        appendix_heading = document.add_paragraph(style="Endnote Heading")
        readable_paragraph(appendix_heading, after=10, font_pt=NOTE_FONT_PT, font=NOTE_FONT)
        set_font(appendix_heading.add_run("빠른 정답"), NOTE_FONT_PT, NOTE_FONT)
        appendix_heading.paragraph_format.keep_with_next = True
        for number, question in enumerate(questions, 1):
            printed = question.get("printedNumber") or str(number)
            answer_paragraphs = append_rich(document, f"{printed}번  {question['answer'] or '정답·풀이 미생성'}", keep=False, font_pt=NOTE_FONT_PT, after=5, font=NOTE_FONT)
            if settings.get("answerMode") == "quick":
                answer_paragraphs[0]._p.insert(1, w_element('bookmarkStart', id=number, name=f'solution_{number}'))
                answer_paragraphs[-1]._p.append(w_element('bookmarkEnd', id=number))
        if settings.get("answerMode") != "quick":
            detail_heading = document.add_paragraph(style="Endnote Heading")
            readable_paragraph(detail_heading, after=10, font_pt=NOTE_FONT_PT, font=NOTE_FONT)
            set_font(detail_heading.add_run("상세 풀이"), NOTE_FONT_PT, NOTE_FONT)
            create_solutions(document, questions, wrap_math=settings.get("wrapSolutionMath", False))
    # Apply the rule to every section, including the separate solution section.
    shared_header = document.sections[0]._sectPr.find(qn("w:headerReference"))
    for section_index, section in enumerate(document.sections):
        columns = section._sectPr.find(qn("w:cols"))
        columns.set(qn("w:num"), "1" if form_pages and section_index < len(form_pages) else "2")
        columns.set(qn("w:sep"), "0" if form_pages and section_index < len(form_pages) else "1")
        columns.set(qn("w:space"), str(round((form_pages[section_index]["gapMm"] if form_pages and section_index<len(form_pages) else 9)*1440/25.4)))
        columns.set(qn("w:equalWidth"), "1")
        if section._sectPr.find(qn("w:headerReference")) is None:
            section._sectPr.insert(0, deepcopy(shared_header))
    for root in (document.element, document.styles.element, document.settings.element):
        canonicalize_properties(root)
    for section in document.sections:
        canonicalize_properties(section.header._element)
        canonicalize_properties(section.footer._element)
    path = Path(output).resolve()
    if path.suffix.lower() != ".docx":
        raise ValueError("출력 확장자는 .docx여야 합니다.")
    path.parent.mkdir(parents=True, exist_ok=True)
    # A failed write must not replace the user's previous successful export.
    with tempfile.NamedTemporaryFile(suffix=".docx", dir=path.parent, delete=False) as temp:
        temp_path = Path(temp.name)
    try:
        document.save(temp_path)
        os.replace(temp_path, path)
    finally:
        if temp_path.exists():
            temp_path.unlink()
    question_pages = sum(math.ceil(p.get("estimatedColumns", 2)/2) for p in pages)
    note_page_estimate = 0 if student else max(1, math.ceil(solution_height / (COL_HEIGHT_PT*2*.9)))
    manifest = {"path": str(path), "pageCountEstimate": question_pages+note_page_estimate,
                "questionPageCountEstimate": question_pages, "solutionPageCountEstimate": note_page_estimate,
                "paginationVerified": False, "warnings": warnings, "layout": pages}
    path.with_suffix(".layout.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def main(argv=None):
    parser = argparse.ArgumentParser(description="문제생성기 DOCX 내보내기")
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args(argv)
    try:
        snapshot = json.loads(Path(args.input).read_text(encoding="utf-8-sig"))
        result = export_document(snapshot, args.output)
        print(json.dumps(result, ensure_ascii=False))
        return 0
    except (ValueError, OSError, KeyError, TypeError, etree.XMLSyntaxError) as error:
        print(f"DOCX 내보내기 실패: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    raise SystemExit(main())
