"""Native document geometry: independent columns and editable midpoint cells."""
import copy, json, os, sys, tempfile, unittest, zipfile
from pathlib import Path
sys.path.insert(0,os.environ.get('EXAM_EXPORT_ENGINE_DIR',str(Path(__file__).resolve().parents[1]/'scripts')))
import export_docx as d, export_hwpx as h
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from lxml import etree as E

class PaperFormLayout(unittest.TestCase):
    def test_native_header_uses_measured_font_without_changing_title(self):
        import paper_form
        doc=Document();title='긴 제목 전체 유지 '*12
        paper_form.configure(doc.sections[0],title=title,title_font_pt=4.5)
        run=doc.sections[0].header.paragraphs[-1].runs[0]
        self.assertEqual(run.text,title);self.assertEqual(run.font.size.pt,4.5)
        self.assertEqual(run._element.rPr.rFonts.get(qn('w:eastAsia')),'맑은 고딕')

    def test_midpoints_keep_native_content_and_first_page_geometry(self):
        with tempfile.TemporaryDirectory() as folder:
            out=Path(folder);background=out/'form.png';Image.new('RGB',(210,297),'white').save(background)
            forms=[dict(topMm=55 if i==0 else 35,bottomMm=22,leftMm=22,rightMm=22,gapMm=9,introMm=42 if i==0 else 0,backgroundPath=str(background),description='first' if i==0 else 'later') for i in range(2)]
            source=dict(title='긴 시험지 제목 배치 검증',paperFormPages=forms,settings=dict(bodyFontSize=10,workspaceLines=0,answerMode='quick',quadrantLayout=True,measuredPages=[[['q0','q1'],['q2','q3']],[['q4','q5'],['q6','q7']]]),questions=[dict(id='q'+str(i),sourceId='q'+str(i),kind='original',body='검증문항 '+str(i)+r' 직선 $\ell$과 $\frac{2}{3}$을 확인하시오.',answer='1',solution='풀이',choices=['1','2','3','4','5']) for i in range(8)])
            before=copy.deepcopy(source);d.export_document(source,out/'exam.docx');self.assertEqual(source,before)
            doc=Document(out/'exam.docx');self.assertEqual(len(doc.sections),3)
            self.assertEqual([s._sectPr.find(qn('w:cols')).get(qn('w:num')) for s in doc.sections],['1','1','2'])
            for i,table in enumerate(doc.tables[:2]):
                self.assertEqual(len(table.columns),3)
                for column in [0,1]:
                    inner=table.cell(0,column*2).tables[0]
                    self.assertEqual(len(inner.rows),2)
                    intro=forms[i]['introMm'] if column==0 else 0
                    expected=((297-forms[i]['topMm']-forms[i]['bottomMm'])+intro)/2*72/25.4
                    self.assertAlmostEqual(inner.rows[0].height.pt,expected,delta=.06)
                    self.assertEqual(inner.rows[0]._tr.trPr.find(qn('w:trHeight')).get(qn('w:hRule')),'atLeast')
                    self.assertIn('검증문항',inner.cell(1,0).text)
            result=h.export_document(source,out/'exam.hwpx')
            self.assertTrue(result['packageVerified'])
            with zipfile.ZipFile(out/'exam.hwpx') as z:
                first=E.fromstring(z.read('Contents/section0.xml'));self.assertGreaterEqual(len(first.findall('.//hp:equation',h.NS)),8)
                # Native Hancom adds header/footer reserves to body margins.
                # Form backgrounds already contain these elements; Word's
                # edge-distance footer must not reduce the measured body area.
                for page in range(len(forms)):
                    section=E.fromstring(z.read(f'Contents/section{page}.xml'))
                    margin=section.find('.//hp:pagePr/hp:margin',h.NS)
                    self.assertEqual(margin.get('header'),'0')
                    self.assertEqual(margin.get('footer'),'0')
                heights=[int(x.get('height')) for x in first.findall('.//hp:cellSz',h.NS)]
                self.assertGreater(max(heights),30000)

    def test_adaptive_browser_positions_reach_native_editable_rows(self):
        import paper_form
        from docx.shared import Mm
        doc=Document()
        nodes=[]
        for text in ['첫 문제', '둘째 문제']:
            p=doc.add_paragraph(text);nodes.append([p._p]);doc._element.body.remove(p._p)
        table=paper_form.question_column(doc,nodes,700,30,220,[30,92])
        self.assertAlmostEqual(table.rows[0].height.mm,92,delta=.02)
        self.assertEqual(table.rows[0]._tr.trPr.find(qn('w:trHeight')).get(qn('w:hRule')),'atLeast')
        self.assertIn('둘째 문제',table.cell(1,0).text)
        with tempfile.TemporaryDirectory() as folder:
            background=Path(folder)/'form.png';Image.new('RGB',(10,10),'white').save(background)
            form=dict(topMm=30,bottomMm=22,leftMm=22,rightMm=22,gapMm=9,introMm=30,backgroundPath=str(background),questionTopsMm=[[30,92],[]])
            snapshot=dict(paperFormPages=[form],settings=dict(measuredPages=[[['a','b'],[]]]))
            self.assertEqual(paper_form.pages(snapshot)[0]['questionTopsMm'],[[30,92],[]])
            for bad in [[[30],[]],[[30,20],[]],[[30,300],[]],[[30,float('nan')],[]]]:
                broken=copy.deepcopy(snapshot);broken['paperFormPages'][0]['questionTopsMm']=bad
                with self.assertRaises(ValueError):paper_form.pages(broken)

if __name__=='__main__':unittest.main()
