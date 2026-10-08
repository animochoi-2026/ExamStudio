"""Cross-format school-math and failure-atomicity contract; no private inputs."""
import copy, json, os, sys, tempfile, unittest, zipfile, subprocess
from pathlib import Path
sys.path.insert(0, os.environ.get('EXAM_EXPORT_ENGINE_DIR',str(Path(__file__).resolve().parents[1]/'scripts')))
import word_math as w
import export_docx as d
import export_hwpx as h
from lxml import etree as E

FORMULAS = [
 r'\ell_1\parallel\ell_2', r'\frac{2x+1}{3}=\sqrt{5}', r'\sqrt[3]{8}=2',
 r'\angle ABC=30^\circ', r'\overline{AB}\perp\overline{CD}',
 r'\overset{\frown}{AB}', r'\vec{a}+\overrightarrow{AB}',
 r'\alpha+\beta=\pi', r'\sin^2 x+\cos^2 x=1',
 r'\log_2 8=3', r'\sum_{k=1}^{n}k', r'\int_0^1 x^2',
 r'\binom{5}{2}=10', r'\left|x-1\right|\leq 2',
 r'A\cap B=\varnothing', r'\mathrm{cm}^2',
 r'\begin{cases}x+1&x<0\\2x&x\geq0\end{cases}',
 r'\begin{pmatrix}1&2\\3&4\end{pmatrix}',
 r'\begin{aligned}x&=1+2\\&=3\end{aligned}',
 r'\boxed{a<0\ \text{또는}\ 0<a<1}',
 r'\frac12\times32=\boxed{16}',
]
def question(i=1, **values):
    return dict(id=f'q{i}',sourceId=f'q{i}',kind='original',body='문제 $x+1=2$',answer='$1$',solution='$x=1$',**values)
def snapshot(questions):
    return {'title':'출력 호환성 검증','questions':questions,'settings':{'bodyFontSize':10,'answerMode':'detailed','quadrantLayout':True,'workspaceLines':0,'measuredPages':[[[q['id']],[]] for q in questions]}}

class ExportContract(unittest.TestCase):
    def test_all_declared_symbols_and_functions_survive_both_native_converters(self):
        for command, symbol in w.SYMBOLS.items():
            with self.subTest(command=command):
                tree=w.parse_latex('\\'+command)
                self.assertEqual(tree.children[0].value,symbol)
                self.assertIn(symbol,''.join(w.latex_to_omml('\\'+command).itertext()))
                self.assertTrue(h.equation(tree))
        for command in w.FUNCTIONS:
            self.assertIn(command,''.join(w.latex_to_omml('\\'+command+' x').itertext()))
            self.assertTrue(h.equation(w.parse_latex('\\'+command+' x')))

    def test_representative_corpus_generates_native_docx_and_hwpx_without_source_changes(self):
        questions=[]
        for i, formula in enumerate(FORMULAS):
            q=question(i);q.update(body='다음 수식을 확인하시오. $'+formula+'$',solution='$'+formula+'$');questions.append(q)
        source=snapshot(questions);original=copy.deepcopy(source)
        with tempfile.TemporaryDirectory() as directory:
            for engine, suffix in [(d,'docx'),(h,'hwpx')]:
                output=Path(directory)/('corpus.'+suffix);engine.export_document(source,output)
                with zipfile.ZipFile(output) as z:
                    self.assertIsNone(z.testzip())
                    if suffix=='docx':
                        xml=E.fromstring(z.read('word/document.xml'))
                        self.assertGreaterEqual(len(xml.findall('.//{'+w.M+'}oMath')),len(FORMULAS)*2)
                    else:
                        self.assertTrue(h.validate_package(output.read_bytes())['packageVerified'])
                        native=json.loads(z.read('Contents/examstudio.json'))
                        for formula in FORMULAS:
                            self.assertTrue(any(formula in (r.get('sourceLatex'),r.get('latex')) for r in native['equations']),formula)
        self.assertEqual(source,original)

    def test_all_bad_questions_reported_together_and_prior_file_preserved(self):
        a,b=question(1),question(2);a['body']=r'$\unknownOne$';b['solution']=r'$\unknownTwo$'
        source=snapshot([a,b])
        with tempfile.TemporaryDirectory() as directory:
            for engine,suffix in [(d,'docx'),(h,'hwpx')]:
                output=Path(directory)/('previous.'+suffix);output.write_bytes(b'previous export')
                with self.assertRaises(w.MathSyntaxError) as raised:engine.export_document(source,output)
                for text in ['1번 본문','2번 상세 풀이','unknownOne','unknownTwo']:self.assertIn(text,str(raised.exception))
                self.assertEqual(output.read_bytes(),b'previous export')

    def test_hwpx_specific_limit_reported_for_all_questions_before_composition(self):
        a,b=question(1),question(2);a['body']=r'$\frac{\boxed{x}}{2}$';b['answer']=r'$x^{\boxed{2}}$'
        with tempfile.TemporaryDirectory() as directory:
            output=Path(directory)/'out.hwpx'
            with self.assertRaises(w.MathSyntaxError) as raised:h.export_document(snapshot([a,b]),output)
            self.assertIn('1번 본문',str(raised.exception));self.assertIn('2번 정답',str(raised.exception));self.assertFalse(output.exists())

    def test_alignment_wrap_preserves_atoms_and_never_cuts_nested_environment(self):
        formula=r'\begin{aligned}\angle ABC+\angle BCD&=180^\circ\\\angle ABC&=180^\circ-60^\circ=120^\circ\end{aligned}'
        rows=d.solution_math_lines(formula,9)
        def atoms(e):
            if e.kind=='seq':return [a for child in e.children for a in atoms(child)]
            if e.kind=='array' and e.value in ('aligned','gathered'):return [a for cells in e.children for cell in cells for a in atoms(cell)]
            return [e]
        self.assertEqual(atoms(w.parse_latex(formula)),[a for row in rows for a in atoms(w.parse_latex(row))])
        self.assertIsNone(d.solution_alignment_rows(r'\begin{aligned}a&=b&c&=d\end{aligned}'))
        with self.assertRaises(w.MathSyntaxError) as raised:d.solution_math_lines(r'\begin{pmatrix}'+('1234567890'*10)+r'&2\\3&4\end{pmatrix}',9)
        self.assertNotIn('end가 없습니다',str(raised.exception))

    def test_stale_three_question_column_cannot_bypass_default_layout(self):
        questions=[question(i) for i in range(3)];source=snapshot(questions)
        source['settings']['measuredPages']=[[[q['id'] for q in questions],[]]]
        with tempfile.TemporaryDirectory() as directory:
            for engine,suffix in [(d,'docx'),(h,'hwpx')]:
                with self.assertRaisesRegex(ValueError,'최대 2문항'):engine.export_document(source,Path(directory)/('out.'+suffix))

    def test_text_solution_profile_preserves_text_without_desktop_only_figure_paths(self):
        q=question();q['solution']='1. 풀이\n$x=1$'
        q['solutionGuide']={'version':1,'basis':'synthetic','steps':[{'id':'step1','title':'풀이','text':'$x=1$','view':{'points':['A','B']}}]}
        source=snapshot([q]);before=copy.deepcopy(source)
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError,'보조그림'):d.export_document(source,Path(directory)/'required.docx')
            source['settings']['solutionGuideMode']='text'
            for engine,suffix in [(d,'docx'),(h,'hwpx')]:engine.export_document(source,Path(directory)/('text.'+suffix))
        self.assertEqual(source['questions'],before['questions'])


class PaperFormContract(unittest.TestCase):
    def test_form_layers_and_free_logo_preserve_native_questions(self):
        import hashlib
        from PIL import Image
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);background=root/'form.png';logo=root/'logo.png'
            Image.new('RGB',(210,297),'white').save(background)
            Image.new('RGBA',(30,20),(150,20,30,102)).save(logo)
            qs=[question(i) for i in range(3)];source=snapshot(qs)
            source['settings'].update(quadrantLayout=False,showStudentNameLine=False,answerMode='quick',solutionGuideMode='text')
            source['paperFormPages']=[dict(backgroundPath=str(background),description='FORM_SCOPE_SCORE_'+str(i),topMm=58 if i==0 else 41,bottomMm=22,leftMm=22,rightMm=22,gapMm=9,introMm=48 if i==0 else 0) for i in range(3)]
            source['paperFormPages'][0]['logo']=dict(path=str(logo),xMm=130,yMm=250,widthMm=30,heightMm=22.5)
            hashes={hashlib.sha256(p.read_bytes()).hexdigest() for p in [background,logo]}
            for engine,suffix in [(d,'docx'),(h,'hwpx')]:
                output=root/('form.'+suffix);engine.export_document(source,output)
                with zipfile.ZipFile(output) as z:
                    native_images={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith(('word/media/','BinData/'))}
                    self.assertTrue(hashes<=native_images)
                    texts='\n'.join(z.read(n).decode('utf8') for n in z.namelist() if n.endswith('.xml'))
                    for i in range(3):self.assertIn('FORM_SCOPE_SCORE_'+str(i),texts)
                    if suffix=='docx':
                        xml=E.fromstring(z.read('word/document.xml'))
                        self.assertGreaterEqual(len(xml.findall('.//{'+w.M+'}oMath')),3)
                        self.assertEqual(len(xml.findall('.//{'+d.W_NS+'}sectPr')),4)
                        headers=[E.fromstring(z.read(n)) for n in z.namelist() if n.startswith('word/header') and n.endswith('.xml')]
                        self.assertTrue(any(n.get('behindDoc')=='0' for x in headers for n in x.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor')))
                    else:
                        self.assertTrue(h.validate_package(output.read_bytes())['packageVerified'])
                        native=json.loads(z.read('Contents/examstudio.json'));self.assertGreaterEqual(len(native['equations']),3)
                        sec=E.fromstring(z.read('Contents/section0.xml'))
                        logo_node=sec.find(".//hp:pic[@textWrap='IN_FRONT_OF_TEXT']",h.NS);self.assertIsNotNone(logo_node)
                        self.assertEqual(int(logo_node.find('hp:pos',h.NS).get('horzOffset')),round(130*7200/25.4))
                        self.assertEqual(len([n for n in z.namelist() if n.startswith('Contents/section') and n.endswith('.xml')]),4)

    def test_invalid_form_page_count_does_not_replace_previous_export(self):
        source=snapshot([question()]);source['paperFormPages']=[]
        with tempfile.TemporaryDirectory() as directory:
            for engine,suffix in [(d,'docx'),(h,'hwpx')]:
                output=Path(directory)/('previous.'+suffix);output.write_bytes(b'previous')
                with self.assertRaises(ValueError):engine.export_document(source,output)
                self.assertEqual(output.read_bytes(),b'previous')


class DeployedWorkerContract(unittest.TestCase):
    def test_bundled_worker_matches_application_and_current_difficulty_contract(self):
        site=Path(os.environ['EXAM_EXPORT_ENGINE_DIR']).resolve().parent
        program=r'''
const fs=require('fs'),vm=require('vm'),assert=require('assert/strict'),path=require('path');
const site=process.argv[1];let message;
const sandbox={structuredClone,performance,self:{postMessage:data=>{message=data;}}};
vm.createContext(sandbox);vm.runInContext(fs.readFileSync(path.join(site,'exam-composition-worker.js'),'utf8'),sandbox);
sandbox.self.onmessage({data:{}});const version=message.version;
assert.ok(version);assert.ok(fs.readFileSync(path.join(site,'app.js'),'utf8').includes(version));
const candidates=Array.from({length:120},(_,i)=>({question_id:'synthetic-'+i,revision_id:'revision-'+i,metadata:{source:{grade:'중2'},content:{responseType:'single_choice'},classification:{primaryUnit:{id:'m2-6.3'},types:[{id:'type-'+i}]},difficulty:{criteriaVersion:'fixed-learner-access-v2',aiScore:i<14?2:i<72?5.5:8.3}}}));
const rules={count:10,units:['m2-6.3'],profile:{low:10,middle:60,high:30,targetAverage:5.7},scopeDistribution:true};
const run=(rows,v=version)=>{sandbox.self.onmessage({data:{version:v,candidates:rows,rules}});return message;};
const result=run(candidates).result;assert.equal(result.eligible,120);assert.equal(result.complete,true);assert.equal(result.items.length,10);
assert.equal(JSON.stringify(result.items.reduce((a,q)=>(a[result.profileSummary.assignment[q.question_id]]++,a),{low:0,middle:0,high:0})),JSON.stringify({low:1,middle:6,high:3}));
const shortage=run(candidates.slice(14)).result;assert.equal(shortage.complete,false);assert.equal(shortage.shortages.find(s=>s.label==='하').missing,1);
assert.match(run(candidates,'stale').error.message,/버전/);
console.log('PASS: deployed application/worker version, 120 eligible, exact quotas, genuine shortage, stale worker');
'''
        result=subprocess.run(['node','-e',program,str(site)],capture_output=True,text=True,encoding='utf-8',timeout=30)
        self.assertEqual(result.returncode,0,result.stdout+result.stderr)

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

class PaperFormLayout(unittest.TestCase):
    def test_native_header_uses_measured_font_without_changing_title(self):
        import paper_form
        doc=Document();title='긴 제목 전체 유지 '*12
        paper_form.configure(doc.sections[0],title=title,title_font_pt=4.5)
        run=doc.sections[0].header.paragraphs[-1].runs[0]
        self.assertEqual(run.text,title);self.assertEqual(run.font.size.pt,4.5)
        self.assertEqual(run._element.rPr.rFonts.get(qn('w:eastAsia')),'맑은 고딕')

    def test_midpoints_keep_native_content_and_first_page_geometry(self):
        with tempfile.TemporaryDirectory() as folder:
            out=Path(folder);background=out/'form.png';Image.new('RGB',(210,297),'white').save(background)
            forms=[dict(topMm=55 if i==0 else 35,bottomMm=22,leftMm=22,rightMm=22,gapMm=9,introMm=42 if i==0 else 0,backgroundPath=str(background),description='first' if i==0 else 'later') for i in range(2)]
            source=dict(title='긴 시험지 제목 배치 검증',paperFormPages=forms,settings=dict(bodyFontSize=10,workspaceLines=0,answerMode='quick',quadrantLayout=True,measuredPages=[[['q0','q1'],['q2','q3']],[['q4','q5'],['q6','q7']]]),questions=[dict(id='q'+str(i),sourceId='q'+str(i),kind='original',body='검증문항 '+str(i)+r' 직선 $\ell$과 $\frac{2}{3}$을 확인하시오.',answer='1',solution='풀이',choices=['1','2','3','4','5']) for i in range(8)])
            before=copy.deepcopy(source);d.export_document(source,out/'exam.docx');self.assertEqual(source,before)
            doc=Document(out/'exam.docx');self.assertEqual(len(doc.sections),3)
            self.assertEqual([s._sectPr.find(qn('w:cols')).get(qn('w:num')) for s in doc.sections],['1','1','2'])
            for i,table in enumerate(doc.tables[:2]):
                self.assertEqual(len(table.columns),3)
                for column in [0,1]:
                    inner=table.cell(0,column*2).tables[0]
                    self.assertEqual(len(inner.rows),2)
                    intro=forms[i]['introMm'] if column==0 else 0
                    expected=((297-forms[i]['topMm']-forms[i]['bottomMm'])+intro)/2*72/25.4
                    self.assertAlmostEqual(inner.rows[0].height.pt,expected,delta=.06)
                    self.assertEqual(inner.rows[0]._tr.trPr.find(qn('w:trHeight')).get(qn('w:hRule')),'atLeast')
                    self.assertIn('검증문항',inner.cell(1,0).text)
            result=h.export_document(source,out/'exam.hwpx')
            self.assertTrue(result['packageVerified'])
            with zipfile.ZipFile(out/'exam.hwpx') as z:
                first=E.fromstring(z.read('Contents/section0.xml'));self.assertGreaterEqual(len(first.findall('.//hp:equation',h.NS)),8)
                heights=[int(x.get('height')) for x in first.findall('.//hp:cellSz',h.NS)]
                self.assertGreater(max(heights),30000)

if __name__=='__main__':unittest.main()
