"""Preserve the editable DOCX and create a separate HWP-compatible image copy.

The source snapshot is mandatory. Every formula must match the existing OMML
and rendered-image manifest, including answer pages, before any output is written.
"""
from __future__ import annotations
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree
from PIL import Image
from word_math import split_math
from export_docx import FONT_PT, NOTE_FONT_PT, formatted_math

M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
IMAGE_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'
PARTS = ('word/document.xml', 'word/endnotes.xml')


def signature(node):
    return node.tag, tuple(sorted(node.attrib.items())), node.text or '', tuple(signature(child) for child in node)


def equation_plan(snapshot):
    questions = [q for q in snapshot['questions'] if q.get('include') is not False]
    parts = {name: [] for name in PARTS}
    def formulas(value, answer=False):
        return [{'latex': content, 'display': kind == 'display' and not answer}
                for kind, content in split_math(value or '') if kind != 'text']
    try:
        body_pt = float(snapshot.get("settings", {}).get("bodyFontSize", 12))
    except (TypeError, ValueError):
        body_pt = 12
    if not 9 <= body_pt <= 18:
        body_pt = 12
    try:
        note_pt = float(snapshot.get("settings", {}).get("solutionFontSize", 9))
    except (TypeError, ValueError):
        note_pt = 9
    if not 9 <= note_pt <= 18:
        note_pt = 9
    def append(value, font_pt, answer=False):
        parts[PARTS[0]].extend({**f, 'fontPt': font_pt} for f in formulas(value, answer))
    for question in questions:
        for value in [question['body'], *question.get('statementBox', []), *question.get('choices', [])]:
            append(value, body_pt)
    for question in questions:
        append(question.get('answer'), note_pt)
    for question in questions:
        append(question.get('answer'), note_pt, answer=True)
        from solution_guide_docx import steps as guide_steps, proof_lines
        ordered = guide_steps(question) if question.get('solutionGuide') else None
        if ordered is None:
            append(question.get('solution'), note_pt)
        else:
            for step in ordered:
                append(step['text'], note_pt)
                for line in proof_lines(question, step): append(line, note_pt)
            core=question['solution'].split('\n\n[서술형 채점기준')[0].split('\n\n[참고 코멘트]')[0]
            append(question['solution'][len(core):], note_pt)
    return parts


def image_run(image_path, record, relation, identifier):
    width, height = (record[name] for name in ('widthPt', 'heightPt'))
    if not all(isinstance(value, (int, float)) and 0 <= value < 2000 for value in (width, height)) or not width or not height:
        raise ValueError('잘못된 HWP 수식 이미지 크기입니다.')
    with Image.open(image_path) as image:
        if image.format != 'PNG' or image.width < 1 or image.height < 1:
            raise ValueError('잘못된 HWP 수식 PNG입니다.')
    document = Document()
    run = document.add_paragraph().add_run()
    run.add_picture(str(image_path), width=Pt(width), height=Pt(height))
    result = deepcopy(run._r)
    properties = OxmlElement('w:rPr')
    properties.append(OxmlElement('w:noProof'))
    # Hanword 2018 crashes while paginating endnote images whose runs contain
    # negative w:position. Keep normal inline-image alignment for this copy.
    result.insert(0, properties)
    result.find('.//' + qn('a:blip')).set(qn('r:embed'), relation)
    for node_name in ('wp:docPr', 'pic:cNvPr'):
        node = result.find('.//' + qn(node_name))
        node.set('id', str(identifier))
        node.set('name', f'Equation {identifier}')
        node.set('descr', record['latex'])
    return result


def make_compatible(input_docx, snapshot_path, manifest_path, output):
    source, output = Path(input_docx).resolve(), Path(output).resolve()
    if source == output:
        raise ValueError('편집 가능한 원본 DOCX와 HWP 호환 복사본의 경로는 달라야 합니다.')
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    snapshot = json.loads(Path(snapshot_path).read_text(encoding='utf-8-sig'))
    plan = equation_plan(snapshot)
    manifest_file = Path(manifest_path).resolve()
    manifest = json.loads(manifest_file.read_text(encoding='utf-8-sig'))
    if manifest.get('version') != 1 or not isinstance(manifest.get('images'), list):
        raise ValueError('잘못된 HWP 수식 이미지 목록입니다.')
    images = manifest['images']
    expected_count = sum(len(values) for values in plan.values())
    if len(images) != expected_count:
        raise ValueError(f'HWP 수식 이미지 수가 원문과 다릅니다 ({len(images)}/{expected_count}).')
    records = {}
    for record in images:
        key = record.get('part'), record.get('index')
        if key in records:
            raise ValueError('중복된 HWP 수식 이미지입니다.')
        records[key] = record
    with zipfile.ZipFile(source) as package:
        blobs = {name: package.read(name) for name in package.namelist()}
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    identifier = 1
    source_pictures = 0
    for name, blob in blobs.items():
        if name.startswith('word/') and name.endswith('.xml'):
            tree = etree.fromstring(blob, parser)
            source_pictures += len(tree.findall('.//' + qn('w:drawing')))
            for node in tree.iter(qn('wp:docPr')):
                identifier = max(identifier, int(node.get('id', '0')) + 1)
    counts = {}
    for part, formulas in plan.items():
        if part not in blobs:
            if formulas:
                raise ValueError('HWP 호환 문서의 수식이 있는 XML 부분을 찾을 수 없습니다.')
            continue
        root = etree.fromstring(blobs[part], parser)
        maths = list(root.iter(f'{{{M}}}oMath'))
        if len(maths) != len(formulas):
            raise ValueError(f'{part}: 원본 DOCX와 수식 원문 개수가 다릅니다 ({len(maths)}/{len(formulas)}).')
        rels_name = str(Path(part).parent.as_posix()) + '/_rels/' + Path(part).name + '.rels'
        rels = etree.fromstring(blobs[rels_name], parser) if rels_name in blobs else etree.Element(f'{{{REL}}}Relationships', nsmap={None: REL})
        used_ids = {node.get('Id') for node in rels}
        for index, (math, formula) in enumerate(zip(maths, formulas)):
            if signature(math) != signature(formatted_math(formula['latex'], font_pt=formula['fontPt'])):
                raise ValueError(f'{part}의 {index+1}번째 수식이 원문과 다릅니다. 먼저 DOCX를 다시 내보내 주세요.')
            record = records.get((part, index))
            if not record or any(record.get(field) != formula[field] for field in ('latex', 'display', 'fontPt')):
                raise ValueError(f'{part}의 {index+1}번째 수식 이미지가 원문과 다릅니다.')
            image_path = (manifest_file.parent / record['file']).resolve()
            if not image_path.is_relative_to(manifest_file.parent) or not image_path.is_file():
                raise ValueError('HWP 수식 이미지가 준비 폴더 밖에 있거나 존재하지 않습니다.')
            media_name = f'word/media/hwp-equation-{identifier:05d}.png'
            while media_name in blobs:
                identifier += 1
                media_name = f'word/media/hwp-equation-{identifier:05d}.png'
            relation = f'rIdHwpEquation{identifier}'
            while relation in used_ids:
                relation += 'x'
            used_ids.add(relation)
            etree.SubElement(rels, f'{{{REL}}}Relationship', Id=relation, Type=IMAGE_REL, Target='media/' + Path(media_name).name)
            run = image_run(image_path, record, relation, identifier)
            parent = math.getparent()
            if parent.tag == f'{{{M}}}oMathPara':
                if len(parent.findall(f'{{{M}}}oMath')) != 1:
                    raise ValueError('하나의 문단에 여러 Office Math 개체가 있습니다.')
                parent.getparent().replace(parent, run)
            else:
                parent.replace(math, run)
            blobs[media_name] = image_path.read_bytes()
            identifier += 1
        if root.find('.//' + f'{{{M}}}oMath') is not None:
            raise ValueError('변환되지 않은 HWP 수식이 남았습니다.')
        counts[part] = len(maths)
        blobs[part] = etree.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)
        blobs[rels_name] = etree.tostring(rels, encoding='UTF-8', xml_declaration=True, standalone=True)
    types = etree.fromstring(blobs['[Content_Types].xml'], parser)
    if not any(node.get('Extension', '').lower() == 'png' for node in types):
        etree.SubElement(types, f'{{{CT}}}Default', Extension='png', ContentType='image/png')
    blobs['[Content_Types].xml'] = etree.tostring(types, encoding='UTF-8', xml_declaration=True, standalone=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix='.docx', delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(temporary_path, 'w', zipfile.ZIP_DEFLATED) as package:
            for name, blob in blobs.items():
                package.writestr(name, blob)
        temporary_path.replace(output)
    finally:
        temporary_path.unlink(missing_ok=True)
    if hashlib.sha256(source.read_bytes()).hexdigest() != source_hash:
        raise ValueError('변환 중 원본 DOCX가 변경되었습니다.')
    return {'ok': True, 'output': str(output), 'equationImages': expected_count,
            'pictureImages': expected_count + source_pictures, 'parts': counts, 'sourceSha256': source_hash}


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stderr.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-docx', required=True)
    parser.add_argument('--snapshot', required=True)
    parser.add_argument('--images', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    try:
        print(json.dumps(make_compatible(args.input_docx, args.snapshot, args.images, args.output), ensure_ascii=False))
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)
